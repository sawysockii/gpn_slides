"""Shape-tree readers: PPTX XML elements → typed IR payloads (spec §10.3–10.7).

Every reader is deterministic Python; no model involvement. Coordinates in the
IR are EMU unless a name says otherwise; affine transforms are computed in Pt.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any, Literal

from lxml import etree

from .assets import AssetStore, sha256_bytes
from .models import (
    Affine2D,
    AxisIR,
    Capability,
    CategoryIR,
    CellStyleIR,
    ChartPayload,
    ChartPointIR,
    ChartSeriesIR,
    ColorExpression,
    ConnectorPayload,
    FieldMetadata,
    GroupPayload,
    HyperlinkIR,
    ImagePayload,
    Issue,
    LineStyleIR,
    MergeRange,
    ObjectIR,
    ParagraphIR,
    PathCommand,
    PlotIR,
    PointPt,
    RectEMU,
    ResolvedTextStyle,
    SemanticMark,
    Severity,
    ShapePayload,
    SourceRef,
    TableCellIR,
    TablePayload,
    TextBreakIR,
    TextPayload,
    TextRunIR,
    TypedValue,
    UnknownPayload,
)
from .units import compose_affine, emu_to_pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
NS = {"a": A, "p": P, "r": R, "c": C}

NSMAP = {"a": A, "p": P, "r": R, "c": C}

_NUM_RE = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")
_QUOTED = re.compile(r'"([^"]*)"|\'([^\']*)\'')

GRAPHIC_DATA_URI_PREFIX = "http://schemas.openxmlformats.org/drawingml/2006/"


def qn(tag: str) -> str:
    prefix, local = tag.split(":", 1)
    return f"{{{NSMAP[prefix]}}}{local}"


def stable_id(*parts: object) -> str:
    """Deterministic content ID from source hash + slide part + shape path + subpath."""
    joined = "|".join(str(p) for p in parts)
    return sha256_bytes(joined.encode("utf-8"))[:24]


@dataclass
class ImportContext:
    """Shared state for one import run (spec §10.3)."""

    graph: Any  # PackageGraph
    asset_store: AssetStore
    source_hash: str
    slide_part: str
    issues: list[Issue] = field(default_factory=list)
    shape_ids_by_slide: dict[str, set[int]] = field(default_factory=dict)
    theme_fonts: dict[str, str] = field(default_factory=dict)
    theme_colors: dict[str, str] = field(default_factory=dict)
    rel_index: dict[tuple[str, str], str] = field(default_factory=dict)

    def rel_target(self, owner_part: str, rel_id: str) -> tuple[str | None, str | None]:
        """Return (resolved_part, external_target) for a relationship id."""
        for rel in self.graph.outgoing.get(owner_part, []):
            if rel.rel_id == rel_id:
                if rel.target_mode == "external":
                    return None, rel.target
                return rel.resolved_part, None
        return None, None

    def add_issue(
        self,
        code: str,
        details: str,
        *,
        rule_id: str = "IMPORT",
        severity: Severity = Severity.WARNING,
        object_ids: list[str] | None = None,
    ) -> None:
        self.issues.append(
            Issue(
                rule_id=rule_id,
                code=code,
                severity=severity,
                slide_id=self.slide_part,
                object_ids=object_ids or [],
                details=details,
                repairable=False,
            )
        )


# ---------------------------------------------------------------------------
# Geometry helpers
# ---------------------------------------------------------------------------

def _emu_int(value: str | None, default: int = 0) -> int:
    if value is None:
        return default
    try:
        return int(float(value))
    except ValueError:
        return default


def _attr(node: etree._Element | None, name: str) -> str | None:
    return None if node is None else node.get(name)


def _find(node: etree._Element, path: str) -> etree._Element | None:
    return node.find(path, namespaces=NS)


def _findall(node: etree._Element, path: str) -> list[etree._Element]:
    return list(node.findall(path, namespaces=NS))


def _xfrm(node: etree._Element | None) -> dict[str, Any] | None:
    """Read an ``a:xfrm``-like element into a dict of EMU fields."""
    if node is None:
        return None
    off = node.find("a:off", namespaces=NS)
    ext = node.find("a:ext", namespaces=NS)
    ch_off = node.find("a:chOff", namespaces=NS)
    ch_ext = node.find("a:chExt", namespaces=NS)
    out: dict[str, Any] = {
        "x": _emu_int(_attr(off, "x")),
        "y": _emu_int(_attr(off, "y")),
        "cx": _emu_int(_attr(ext, "cx")),
        "cy": _emu_int(_attr(ext, "cy")),
        "rot": int(node.get("rot") or 0) / 60000.0,
        "flipH": node.get("flipH") == "1",
        "flipV": node.get("flipV") == "1",
    }
    if ch_off is not None:
        out["chX"] = _emu_int(_attr(ch_off, "x"))
        out["chY"] = _emu_int(_attr(ch_off, "y"))
    if ch_ext is not None:
        out["chCX"] = _emu_int(_attr(ch_ext, "cx"))
        out["chCY"] = _emu_int(_attr(ch_ext, "cy"))
    return out


def _rotation_flip_matrix(box: dict[str, Any]) -> Affine2D:
    """Rotation/flip about the box centre, in Pt (spec §10.3 step 3)."""
    cx = emu_to_pt(box["x"] + box["cx"] / 2)
    cy = emu_to_pt(box["y"] + box["cy"] / 2)
    import math

    angle = math.radians(box.get("rot", 0.0))
    cos_a, sin_a = math.cos(angle), math.sin(angle)
    # T(cx,cy) @ R @ F @ T(-cx,-cy)
    f_a = -1.0 if box.get("flipH") else 1.0
    f_d = -1.0 if box.get("flipV") else 1.0
    r = Affine2D(a=cos_a, b=sin_a, c=-sin_a, d=cos_a, tx=0.0, ty=0.0)
    f = Affine2D(a=f_a, d=f_d)
    t_in = Affine2D(tx=-cx, ty=-cy)
    t_out = Affine2D(tx=cx, ty=cy)
    return compose_affine(compose_affine(t_out, compose_affine(r, f)), t_in)


def group_child_matrix(box: dict[str, Any]) -> Affine2D:
    """Child-space → parent-space matrix for a group (shift, scale, then off)."""
    sx = 1.0
    sy = 1.0
    if box.get("chCX"):
        sx = box["cx"] / box["chCX"]
    if box.get("chCY"):
        sy = box["cy"] / box["chCY"]
    scale = Affine2D(a=sx, d=sy)
    shift = Affine2D(tx=-emu_to_pt(box.get("chX", 0)), ty=-emu_to_pt(box.get("chY", 0)))
    move = Affine2D(tx=emu_to_pt(box["x"]), ty=emu_to_pt(box["y"]))
    local = compose_affine(compose_affine(move, scale), shift)
    outer = _rotation_flip_matrix(box)
    return compose_affine(outer, local)


# ---------------------------------------------------------------------------
# Text reading (§10.4)
# ---------------------------------------------------------------------------

def _color_expression(node: etree._Element | None) -> ColorExpression | None:
    if node is None:
        return None
    srgb = node.find("a:srgbClr", namespaces=NS)
    if srgb is not None:
        return ColorExpression(kind="srgb", value=(srgb.get("val") or "").upper())
    scheme = node.find("a:schemeClr", namespaces=NS)
    if scheme is not None:
        return ColorExpression(kind="scheme", value=scheme.get("val") or "")
    system = node.find("a:systemClr", namespaces=NS)
    if system is not None:
        return ColorExpression(kind="system", value=system.get("val") or "")
    return None


def _fill_color(node: etree._Element | None) -> ColorExpression | None:
    """Read a solidFill colour inside ``node`` (or node itself being the fill)."""
    if node is None:
        return None
    if node.tag == qn("a:solidFill"):
        return _color_expression(node)
    solid = node.find("a:solidFill", namespaces=NS)
    return _color_expression(solid)


def _run_style(rpr: etree._Element | None, inherited: list[str]) -> ResolvedTextStyle:
    if rpr is None:
        return ResolvedTextStyle(inherited_from=list(inherited))
    typeface = None
    latin = rpr.find("a:latin", namespaces=NS)
    if latin is not None:
        typeface = latin.get("typeface")
    sz = rpr.get("sz")
    size_pt = int(sz) / 100.0 if sz and sz.lstrip("-").isdigit() else None
    bold = rpr.get("b")
    italic = rpr.get("i")
    color = _fill_color(rpr)
    lang = rpr.get("lang")
    scale = rpr.find("a:scale", namespaces=NS)
    font_scale = (
        float(scale.get("val")) / 100000.0
        if scale is not None and scale.get("val")
        else 1.0
    )
    return ResolvedTextStyle(
        typeface=typeface,
        size_pt=size_pt,
        bold=None if bold is None else bold == "1",
        italic=None if italic is None else italic == "1",
        underline=rpr.get("u"),
        color=color,
        language=lang,
        inherited_from=list(inherited),
        effective_font_scale=font_scale,
    )


def _semantic_marks(rpr: etree._Element | None) -> list[Any]:
    """Baseline/vertical-underline decorations that carry semantics (§7.2)."""
    if rpr is None:
        return []
    marks: list[Any] = []
    baseline = rpr.get("baseline")
    if baseline:
        try:
            value = int(baseline)
        except ValueError:
            value = 0
        if value:
            kind = "superscript" if value > 0 else "subscript"
            marks.append(
                SemanticMark(kind=kind, start=0, end=0, evidence_ref=f"baseline={value}")
            )
    return marks


def _paragraph_props(p: etree._Element) -> dict[str, Any]:
    ppr = p.find("a:pPr", namespaces=NS)
    out: dict[str, Any] = {
        "level": 0,
        "marL": None,
        "indent": None,
        "list_kind": "none",
        "space_before_pt": None,
        "space_after_pt": None,
        "numbering": None,
        "algn": None,
        "has_ppr": ppr is not None,
        "ppr": ppr,
    }
    if ppr is None:
        return out
    out["level"] = int(ppr.get("lvl") or 0)
    mar = ppr.get("marL")
    ind = ppr.get("indent")
    out["marL"] = int(mar) if mar and mar.lstrip("-").isdigit() else None
    out["indent"] = int(ind) if ind and ind.lstrip("-").isdigit() else None
    out["algn"] = ppr.get("algn")
    spc_bef = ppr.find("a:spcBef/a:spcPts", namespaces=NS)
    spc_aft = ppr.find("a:spcAft/a:spcPts", namespaces=NS)
    if spc_bef is not None and spc_bef.get("val"):
        out["space_before_pt"] = int(spc_bef.get("val") or 0) / 100.0
    if spc_aft is not None and spc_aft.get("val"):
        out["space_after_pt"] = int(spc_aft.get("val") or 0) / 100.0
    if ppr.find("a:buChar", namespaces=NS) is not None:
        out["list_kind"] = "bullet"
    elif ppr.find("a:buAutoNum", namespaces=NS) is not None:
        out["list_kind"] = "numbered"
        auto = ppr.find("a:buAutoNum", namespaces=NS)
        out["numbering"] = {
            "scheme": (auto.get("type") or "arabicPeriod") if auto is not None else "arabicPeriod",
            "start_at": int(ppr.get("startAt") or 1),
        }
    return out


def _hyperlink_for(
    node: etree._Element | None, ctx: ImportContext, fallback_ids: list[str]
) -> HyperlinkIR | None:
    if node is None:
        return None
    # Click is primary; hover-only runs keep their hover link so the IR never
    # silently drops a hyperlink event (full click+hover coverage lives in
    # iter_hyperlink_occurrences over raw XML).
    click = node.find("a:hlinkClick", namespaces=NS)
    hover = node.find("a:hlinkHover", namespaces=NS)
    chosen = click if click is not None else hover
    if chosen is None:
        return None
    rid = chosen.get(qn("r:id")) or ""
    tooltip = chosen.get("tooltip")
    part, external = ctx.rel_target(ctx.slide_part, rid) if rid else (None, None)
    if external:
        return HyperlinkIR(
            target_kind="external", uri=external, tooltip=tooltip, source_rel_id=rid
        )
    if part:
        return HyperlinkIR(
            target_kind="internal_slide",
            target_slide_ref=part,
            tooltip=tooltip,
            source_rel_id=rid,
        )
    return HyperlinkIR(target_kind="internal_slide", uri=None, source_rel_id=rid, tooltip=tooltip)


def read_text_body(
    tx_body: etree._Element, ctx: ImportContext, subpath_root: str, shape_key: object
) -> TextPayload:
    """Read ``a:txBody`` / ``p:txBody`` into a TextPayload (spec §10.4)."""
    paragraphs: list[ParagraphIR] = []
    body_pr = tx_body.find("a:bodyPr", namespaces=NS)
    margins = (
        _emu_int(_attr(body_pr, "lIns"), 91440),
        _emu_int(_attr(body_pr, "tIns"), 45720),
        _emu_int(_attr(body_pr, "rIns"), 91440),
        _emu_int(_attr(body_pr, "bIns"), 45720),
    )
    anchor = _attr(body_pr, "anchor") or "top"
    direction = "horz"
    if body_pr is not None:
        direction = "vert" if body_pr.get("vert") else "horz"

    role = _guess_text_role(tx_body)

    for p_index, p in enumerate(tx_body.findall("a:p", namespaces=NS)):
        props = _paragraph_props(p)
        para_id = stable_id(
            ctx.source_hash, ctx.slide_part, shape_key, f"{subpath_root}/p{p_index}"
        )
        runs: list[TextRunIR] = []
        breaks: list[TextBreakIR] = []
        inherited = ["paragraph", "shape"] + (["list"] if props["has_ppr"] else [])

        r_index = 0
        for child in p:
            if child.tag == qn("a:r"):
                rpr = child.find("a:rPr", namespaces=NS)
                text_node = child.find("a:t", namespaces=NS)
                text = text_node.text or "" if text_node is not None else ""
                run_id = f"{para_id}/r{r_index}"
                r_index += 1
                runs.append(
                    TextRunIR(
                        id=run_id,
                        source_ref=SourceRef(
                            deck_sha256=ctx.source_hash,
                            slide_part=ctx.slide_part,
                            shape_id=_shape_id_of(shape_key),
                            group_path=_group_path_of(shape_key),
                            subpath=f"{subpath_root}/p{p_index}/r{len(runs) - 1}",
                        ),
                        text=text,
                        source_style=_run_style(rpr, inherited),
                        semantic_marks=_semantic_marks(rpr),
                        hyperlink=_hyperlink_for(rpr, ctx, [run_id]),
                    )
                )
            elif child.tag == qn("a:br"):
                rpr = child.find("a:rPr", namespaces=NS)
                run_id = f"{para_id}/br{len(breaks)}"
                breaks.append(TextBreakIR(after_run_id=run_id, offset=r_index, kind="soft"))
                # Preserve the break as an empty run carrier with its style.
                runs.append(
                    TextRunIR(
                        id=run_id,
                        source_ref=SourceRef(
                            deck_sha256=ctx.source_hash,
                            slide_part=ctx.slide_part,
                            shape_id=_shape_id_of(shape_key),
                            group_path=_group_path_of(shape_key),
                            subpath=f"{subpath_root}/p{p_index}/br{len(breaks) - 1}",
                        ),
                        text="\n",
                        source_style=_run_style(rpr, inherited),
                    )
                )
                r_index += 1
            elif child.tag == qn("a:fld"):
                fld_type = child.get("id") or child.get("type") or ""
                text_node = child.find("a:t", namespaces=NS)
                text = text_node.text or "" if text_node is not None else ""
                rpr = child.find("a:rPr", namespaces=NS)
                run_id = f"{para_id}/f{r_index}"
                r_index += 1
                field_meta = FieldMetadata(
                    field_id=fld_type,
                    field_type=child.get("type") or "",
                    cached_text=text,
                    source_ref=SourceRef(
                        deck_sha256=ctx.source_hash,
                        slide_part=ctx.slide_part,
                        shape_id=_shape_id_of(shape_key),
                        group_path=_group_path_of(shape_key),
                        subpath=f"{subpath_root}/p{p_index}/fld{fld_type}",
                    ),
                    raw_xml_asset=None,
                )
                run = TextRunIR(
                    id=run_id,
                    source_ref=SourceRef(
                        deck_sha256=ctx.source_hash,
                        slide_part=ctx.slide_part,
                        shape_id=_shape_id_of(shape_key),
                        group_path=_group_path_of(shape_key),
                        subpath=f"{subpath_root}/p{p_index}/fld{fld_type}",
                    ),
                    text=text,
                    source_style=_run_style(rpr, inherited),
                    hyperlink=_hyperlink_for(rpr, ctx, [run_id]),
                    field=field_meta,
                    kind="field",
                )
                runs.append(run)

        numbering = None
        if props["numbering"]:
            numbering = props["numbering"]
        paragraphs.append(
            ParagraphIR(
                id=para_id,
                source_ref=SourceRef(
                    deck_sha256=ctx.source_hash,
                    slide_part=ctx.slide_part,
                    shape_id=_shape_id_of(shape_key),
                    group_path=_group_path_of(shape_key),
                    subpath=f"{subpath_root}/p{p_index}",
                ),
                runs=runs,
                breaks=breaks,
                list_kind=props["list_kind"],
                level=props["level"],
                numbering=numbering,  # type: ignore[arg-type]
                source_indent_emu=props["marL"],
                source_hanging_emu=props["indent"],
                space_before_pt=props["space_before_pt"],
                space_after_pt=props["space_after_pt"],
            )
        )

    return TextPayload(
        paragraphs=paragraphs,
        role=role,
        text_direction=direction,
        vertical_anchor=anchor,
        internal_margins_emu=margins,
    )


def _guess_text_role(tx_body: etree._Element) -> str:
    body_pr = tx_body.find("a:bodyPr", namespaces=NS)
    if body_pr is not None and body_pr.find("a:normAutofit", namespaces=NS) is not None:
        return "body"
    return "unknown"


def _shape_id_of(shape_key: object) -> int:
    if isinstance(shape_key, tuple):
        return int(shape_key[0])
    if isinstance(shape_key, int):
        return shape_key
    return 0


def _group_path_of(shape_key: object) -> list[int]:
    if isinstance(shape_key, tuple) and len(shape_key) > 1:
        path = shape_key[1]
        if isinstance(path, (list, tuple)):
            return [int(x) for x in path]
    return []


# ---------------------------------------------------------------------------
# Table reading (§10.5)
# ---------------------------------------------------------------------------

def _cell_paragraphs(
    tc: etree._Element, ctx: ImportContext, subpath: str, shape_key: object
) -> list[ParagraphIR]:
    tx = tc.find("a:txBody", namespaces=NS)
    if tx is None:
        return []
    payload = read_text_body(tx, ctx, subpath, shape_key)
    return payload.paragraphs


def _typed_value_from_paragraphs(paragraphs: list[ParagraphIR]) -> TypedValue | None:
    text = "".join(run.text for p in paragraphs for run in p.runs).strip()
    if text == "":
        return TypedValue(state="missing", display_text="")
    if _NUM_RE.match(text):
        return TypedValue(state="present", decimal=_normalize_decimal(text), display_text=text)
    return None


def _normalize_decimal(text: str) -> str:
    try:
        from decimal import Decimal

        return str(Decimal(text))
    except Exception:  # noqa: BLE001 - fall back to raw text
        return text


def read_table(frame: etree._Element, ctx: ImportContext, shape_key: object) -> TablePayload:
    """Read ``a:tbl`` including empty/spanned cells and merge ranges (spec §10.5)."""
    tbl = frame.find(".//a:tbl", namespaces=NS)
    if tbl is None:
        raise ValueError("graphicFrame has no a:tbl")
    grid = tbl.find("a:tblGrid", namespaces=NS)
    col_widths: list[int] = []
    if grid is not None:
        for gc in grid.findall("a:gridCol", namespaces=NS):
            col_widths.append(_emu_int(gc.get("w")))
    cols = max(len(col_widths), 1)

    rows_nodes = tbl.findall("a:tr", namespaces=NS)
    row_heights = [_emu_int(tr.get("h")) for tr in rows_nodes]
    rows = len(rows_nodes)

    # Map each XML cell onto the physical grid using gridSpan/rowSpan occupancy.
    # Exactly one TableCellIR per grid coordinate: covered coordinates get a
    # single spanned placeholder; repeated continuation elements for an
    # already-emitted coordinate are skipped (never re-emitted).
    occupied: dict[tuple[int, int], str] = {}
    emitted: set[tuple[int, int]] = set()
    cells: list[TableCellIR] = []
    merges: list[MergeRange] = []

    for r, tr in enumerate(rows_nodes):
        c = 0
        for tc in tr.findall("a:tc", namespaces=NS):
            span = int(tc.get("gridSpan") or 1)
            vspan = int(tc.get("rowSpan") or 1)
            h_merge = tc.get("hMerge") == "1"
            v_merge = tc.get("vMerge") == "1"

            if h_merge:
                # Covered column of a preceding horizontal span sits at c-1.
                if (r, c - 1) in occupied:
                    col = c - 1
                    advance = 0
                else:
                    col = c
                    advance = 1
                if (r, col) in emitted:
                    # Same grid coordinate already has its cell: the covered
                    # continuation element carries no independent cell.
                    c += 1 if advance else 0
                    continue
            elif v_merge:
                col = c
                advance = 1
                if (r, col) in emitted:
                    c += advance
                    continue
            else:
                while (r, c) in occupied:
                    c += 1
                col = c
                advance = span

            cell_id = stable_id(
                ctx.source_hash, ctx.slide_part, shape_key, f"table/r{r}/c{col}"
            )
            subpath = f"table/r{r}/c{col}"

            if h_merge or v_merge:
                occupied[(r, col)] = "spanned"
                emitted.add((r, col))
                spanned_paragraphs = _cell_paragraphs(tc, ctx, subpath, shape_key)
                cells.append(
                    TableCellIR(
                        id=cell_id,
                        source_ref=SourceRef(
                            deck_sha256=ctx.source_hash,
                            slide_part=ctx.slide_part,
                            shape_id=_shape_id_of(shape_key),
                            group_path=_group_path_of(shape_key),
                            subpath=subpath,
                        ),
                        row=r,
                        col=col,
                        is_spanned=True,
                        paragraphs=spanned_paragraphs,
                        typed_value=_typed_value_from_paragraphs(spanned_paragraphs),
                        visible=False,
                    )
                )
                c += advance
                continue

            for dr in range(vspan):
                for dc in range(span):
                    occupied[(r + dr, col + dc)] = "occupied"

            style = CellStyleIR()
            tc_pr = tc.find("a:tcPr", namespaces=NS)
            if tc_pr is not None:
                fill = _fill_color(tc_pr)
                if fill is not None:
                    style.fill = fill
                anchor = tc_pr.get("anchor")
                if anchor:
                    style.vertical_anchor = anchor
                margins = tuple(
                    _emu_int(tc_pr.get(k), d)
                    for k, d in (
                        ("marL", 45720), ("marT", 45720),
                        ("marR", 45720), ("marB", 45720),
                    )
                )
                style.margins_emu = margins  # type: ignore[assignment]
            paragraphs = _cell_paragraphs(tc, ctx, subpath, shape_key)
            is_origin = span > 1 or vspan > 1
            emitted.add((r, col))
            cells.append(
                TableCellIR(
                    id=cell_id,
                    source_ref=SourceRef(
                        deck_sha256=ctx.source_hash,
                        slide_part=ctx.slide_part,
                        shape_id=_shape_id_of(shape_key),
                        group_path=_group_path_of(shape_key),
                        subpath=subpath,
                    ),
                    row=r,
                    col=col,
                    row_span=vspan,
                    col_span=span,
                    is_merge_origin=is_origin,
                    is_spanned=False,
                    paragraphs=paragraphs,
                    typed_value=_typed_value_from_paragraphs(paragraphs),
                    source_style=style,
                    visible=True,
                )
            )
            if is_origin:
                merges.append(
                    MergeRange(
                        r0=r, c0=col,
                        r1=min(r + vspan - 1, rows - 1),
                        c1=min(col + span - 1, cols - 1),
                    )
                )
            c += advance

        # Complete any grid cells the XML never provided.
        col = 0
        while col < cols:
            if (r, col) not in occupied:
                occupied[(r, col)] = "occupied"
                emitted.add((r, col))
                cell_id = stable_id(
                    ctx.source_hash, ctx.slide_part, shape_key, f"table/r{r}/c{col}"
                )
                cells.append(
                    TableCellIR(
                        id=cell_id,
                        source_ref=SourceRef(
                            deck_sha256=ctx.source_hash,
                            slide_part=ctx.slide_part,
                            shape_id=_shape_id_of(shape_key),
                            group_path=_group_path_of(shape_key),
                            subpath=f"table/r{r}/c{col}",
                        ),
                        row=r,
                        col=col,
                        paragraphs=[],
                        visible=True,
                    )
                )
            col += 1

    if rows and cols and len(cells) < rows * cols:
        ctx.add_issue(
            "TABLE_GRID_INCOMPLETE",
            f"expected {rows * cols} cells, parsed {len(cells)}",
            object_ids=[str(_shape_id_of(shape_key))],
        )

    return TablePayload(
        rows=rows,
        cols=cols,
        cells=cells,
        merges=_dedupe_merges(merges),
        column_widths_emu=col_widths,
        row_heights_emu=row_heights,
    )


def cell_id_key(r: int, c: int) -> str:
    return f"r{r}c{c}"


def _dedupe_merges(merges: list[MergeRange]) -> list[MergeRange]:
    seen: set[tuple[int, int, int, int]] = set()
    out: list[MergeRange] = []
    for m in merges:
        key = (m.r0, m.c0, m.r1, m.c1)
        if key not in seen:
            seen.add(key)
            out.append(m)
    return out


# ---------------------------------------------------------------------------
# Chart reading (§10.6)
# ---------------------------------------------------------------------------

_REF_FORMULA_RE = re.compile(
    r"^'?(?P<sheet>[^'!]+)'?!(?P<range>\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)$"
)


def _ref_values(node: etree._Element | None) -> tuple[str | None, list[tuple[int, str]]]:
    """Read strRef/numRef/strLit/numLit contents → (formula, [(idx, raw)])."""
    if node is None:
        return None, []
    for tag in ("c:strRef", "c:numRef"):
        ref = node.find(tag, namespaces=NS)
        if ref is None:
            continue
        formula = ref.findtext("c:f", default=None, namespaces=NS)
        cache = ref.find("c:strCache" if tag == "c:strRef" else "c:numCache", namespaces=NS)
        values: list[tuple[int, str]] = []
        if cache is not None:
            for pt in cache.findall("c:pt", namespaces=NS):
                idx = int(pt.get("idx") or 0)
                val = pt.findtext("c:v", default="", namespaces=NS) or ""
                values.append((idx, val))
        return formula, values
    for tag in ("c:strLit", "c:numLit"):
        lit = node.find(tag, namespaces=NS)
        if lit is None:
            continue
        values = []
        for pt in lit.findall("c:pt", namespaces=NS):
            idx = int(pt.get("idx") or 0)
            val = pt.findtext("c:v", default="", namespaces=NS) or ""
            values.append((idx, val))
        return None, values
    return None, []


def _series_name(ser: etree._Element) -> tuple[str, str | None]:
    tx = ser.find("c:tx", namespaces=NS)
    if tx is None:
        return "", None
    formula, values = _ref_values(tx)
    if values:
        return values[0][1], formula
    return "", formula


def _typed(raw: str) -> TypedValue | None:
    text = raw.strip()
    if text == "":
        return TypedValue(state="missing", display_text="")
    if _NUM_RE.match(text):
        return TypedValue(state="present", decimal=_normalize_decimal(text), display_text=text)
    return TypedValue(state="unknown", display_text=text)


def _points_from_refs(
    ctx: ImportContext,
    owner: str,
    ser_index: int,
    sub_root: str,
    cat_formula: str | None,
    cats: list[tuple[int, str]],
    val_formula: str | None,
    vals: list[tuple[int, str]],
    *,
    x_vals: list[tuple[int, str]] | None = None,
    y_vals: list[tuple[int, str]] | None = None,
    bubble_vals: list[tuple[int, str]] | None = None,
) -> list[ChartPointIR]:
    count = max(
        len(vals), len(cats), len(x_vals or []), len(y_vals or []), len(bubble_vals or [])
    )
    points: list[ChartPointIR] = []
    cat_map = dict(cats)
    val_map = dict(vals)
    for idx in range(count):
        pid = stable_id(ctx.source_hash, ctx.slide_part, owner, f"{sub_root}/point{idx}")
        category_path = [cat_map[idx]] if idx in cat_map else None
        in_value_range = idx < max(len(vals), len(cats))
        point = ChartPointIR(
            id=pid,
            index=idx,
            category_path=category_path,
            value=(
                _typed(val_map[idx])
                if idx in val_map
                else (TypedValue(state="missing", display_text="") if in_value_range else None)
            ),
            x=_typed(dict(x_vals or {})[idx]) if x_vals and idx in dict(x_vals) else None,
            y=_typed(dict(y_vals or {})[idx]) if y_vals and idx in dict(y_vals) else None,
            bubble_size=(
                _typed(dict(bubble_vals or {})[idx])
                if bubble_vals and idx in dict(bubble_vals)
                else None
            ),
        )
        points.append(point)
    return points


def _axis_list(chart_space: etree._Element, ctx: ImportContext) -> list[AxisIR]:
    axes: list[AxisIR] = []
    plot_area = chart_space.find("c:chart/c:plotArea", namespaces=NS)
    if plot_area is None:
        return axes
    for node in plot_area:
        local = etree.QName(node).localname
        if local not in ("catAx", "valAx", "dateAx", "serAx"):
            continue
        ax_id = node.findtext("c:axId", default="", namespaces=NS) or node.get("val") or ""
        type_map = {"catAx": "category", "valAx": "value", "dateAx": "date", "serAx": "series"}
        scaling = node.find("c:scaling", namespaces=NS)
        log_base = None
        reversed_ = False
        minimum = None
        maximum = None
        if scaling is not None:
            log_base = scaling.findtext("c:logBase", default=None, namespaces=NS)
            orientation = scaling.findtext("c:orientation", default=None, namespaces=NS)
            reversed_ = orientation == "maxMin"
            minimum = scaling.findtext("c:min", default=None, namespaces=NS)
            maximum = scaling.findtext("c:max", default=None, namespaces=NS)
        crosses = node.findtext("c:crosses", default=None, namespaces=NS)
        num_fmt = node.find("c:numFmt", namespaces=NS)
        number_format = num_fmt.get("formatCode") if num_fmt is not None else None
        title_text = None
        title = node.find("c:title", namespaces=NS)
        if title is not None:
            texts = title.findall(".//a:t", namespaces=NS)
            joined = "".join(t.text or "" for t in texts)
            if joined:
                title_text = TextPayload(
                    paragraphs=[
                        ParagraphIR(
                            id=stable_id(ctx.source_hash, ax_id, "title"),
                            source_ref=SourceRef(
                                deck_sha256=ctx.source_hash,
                                slide_part=ctx.slide_part,
                                shape_id=0,
                                subpath=f"chart/axis/{ax_id}/title",
                            ),
                            runs=[],
                        )
                    ],
                    role="caption",
                )
        axes.append(
            AxisIR(
                id=f"ax{ax_id}",
                type=type_map.get(local, "value"),  # type: ignore[arg-type]
                unit_ref=None,
                minimum=minimum,
                maximum=maximum,
                log_base=log_base,
                reversed=reversed_,
                crosses=crosses,
                number_format=number_format,
                title=title_text,
            )
        )
    return axes


ctx_hash_placeholder = ""


def read_chart(frame: etree._Element, ctx: ImportContext, shape_key: object) -> ChartPayload:
    """Read a native chart part: plots, series, points, axes, workbook (§10.6)."""
    chart_ref = frame.find(".//c:chart", namespaces=NS)
    if chart_ref is None:
        raise ValueError("graphicFrame has no c:chart reference")
    rid = chart_ref.get(qn("r:id")) or ""
    chart_part, external = ctx.rel_target(ctx.slide_part, rid)
    if chart_part is None and external is not None:
        ctx.add_issue(
            "CHART_EXTERNAL_WORKBOOK_URI",
            f"chart relationship points outside package: {external}",
            severity=Severity.WARNING,
        )
    if chart_part is None or chart_part not in ctx.graph.parts_by_name:
        ctx.add_issue(
            "CHART_PART_MISSING", f"chart part not found for rId={rid}",
            severity=Severity.ERROR,
        )
        return ChartPayload(
            chart_part="",
            chart_type="unknown",
            data_source_status="missing",
            part_graph_root="",
        )

    xml_bytes = ctx.asset_store.read_bytes(ctx.graph.parts_by_name[chart_part].asset)  # type: ignore[arg-type]
    chart_space = etree.fromstring(
        xml_bytes, parser=etree.XMLParser(resolve_entities=False, no_network=True)
    )
    is_chart_ex = etree.QName(chart_space).localname.lower().startswith("chartex")
    if is_chart_ex:
        ctx.add_issue(
            "CHARTEX_UNSUPPORTED",
            f"chartEx part detected and preserved only as raw XML: {chart_part}",
            severity=Severity.WARNING,
        )

    plots: list[PlotIR] = []
    series: list[ChartSeriesIR] = []
    categories: list[CategoryIR] = []
    plot_area = chart_space.find("c:chart/c:plotArea", namespaces=NS)

    cat_index: dict[str, int] = {}
    if plot_area is not None:
        for plot_node in plot_area:
            local = etree.QName(plot_node).localname
            if not local.endswith("Chart"):
                continue
            ax_ids = [
                f"ax{n.get('val')}"
                for n in plot_node.findall("c:axId", namespaces=NS)
            ]
            plot_id = stable_id(chart_part, local, len(plots), *ax_ids)
            plot_series_ids: list[str] = []
            grouping = plot_node.findtext("c:grouping", default=None, namespaces=NS)
            stacking = plot_node.findtext("c:grouping", default=None, namespaces=NS)
            for s_index, ser in enumerate(plot_node.findall("c:ser", namespaces=NS)):
                name, name_formula = _series_name(ser)
                idx = int(ser.findtext("c:idx", default=str(s_index), namespaces=NS) or s_index)
                cat_node = ser.find("c:cat", namespaces=NS)
                val_node = ser.find("c:val", namespaces=NS)
                x_node = ser.find("c:xVal", namespaces=NS)
                y_node = ser.find("c:yVal", namespaces=NS)
                b_node = ser.find("c:bubbleSize", namespaces=NS)
                cat_formula, cats = _ref_values(cat_node)
                val_formula, vals = _ref_values(val_node)
                x_vals = _ref_values(x_node)[1] if x_node is not None else None
                y_vals = _ref_values(y_node)[1] if y_node is not None else None
                b_vals = _ref_values(b_node)[1] if b_node is not None else None

                for c_idx, c_raw in cats:
                    key = f"{cat_formula or 'cat'}#{c_idx}"
                    if key not in cat_index:
                        cat_index[key] = len(categories)
                        categories.append(
                            CategoryIR(
                                id=stable_id(chart_part, "cat", key),
                                index=cat_index[key],
                                path=[c_raw],
                                raw_value=c_raw,
                            )
                        )
                pts = _points_from_refs(
                    ctx,
                    chart_part,
                    idx,
                    # Plot-scoped subpath: per-plot series positions repeat
                    # across plots, so bare series ids would collide.
                    f"plot{len(plots)}/series{s_index}",
                    cat_formula,
                    cats,
                    val_formula,
                    vals,
                    x_vals=x_vals,
                    y_vals=y_vals,
                    bubble_vals=b_vals,
                )
                sid = stable_id(chart_part, "ser", idx, name)
                formula_refs = [f for f in (cat_formula, val_formula, name_formula) if f]
                series.append(
                    ChartSeriesIR(
                        id=sid,
                        name=name,
                        source_index=idx,
                        points=pts,
                        plot_id=plot_id,
                        axis_group="primary" if len(ax_ids) <= 1 else "primary/secondary",
                        formula_refs=formula_refs,
                    )
                )
                plot_series_ids.append(sid)
            plots.append(
                PlotIR(
                    id=plot_id,
                    type_token=local,
                    series_ids=plot_series_ids,
                    axis_ids=ax_ids,
                    grouping=grouping,
                    stacking=stacking,
                )
            )

    axes = _axis_list(chart_space, ctx)

    # Workbook / data source status.
    workbook_ref: Any = None
    data_status: Literal["embedded", "cache_only", "external", "missing"] = "missing"
    external_uri = None
    ext_data = chart_space.find("c:externalData", namespaces=NS)
    if ext_data is not None:
        rid2 = ext_data.get(qn("r:id")) or ""
        wb_part, wb_external = ctx.rel_target(chart_part, rid2)
        if wb_external:
            data_status = "external"
            external_uri = wb_external
        elif wb_part and wb_part in ctx.graph.parts_by_name:
            wb_bytes = ctx.asset_store.read_bytes(ctx.graph.parts_by_name[wb_part].asset)  # type: ignore[arg-type]
            workbook_ref = ctx.asset_store.put(
                wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            data_status = "embedded"
    if data_status == "missing" and series and any(p.value for s in series for p in s.points):
        data_status = "cache_only"

    chart_xml_asset = ctx.asset_store.put(xml_bytes, "application/xml")
    disp = chart_space.find("c:chart", namespaces=NS)
    display_blanks = disp.get("dispBlanksAs") if disp is not None else "gap"
    chart_type = plots[0].type_token if plots else "unknown"

    if workbook_ref is not None:
        _verify_workbook(ctx, workbook_ref, series, chart_part)

    return ChartPayload(
        chart_part=chart_part,
        chart_type=chart_type,
        plots=plots,
        series=series,
        axes=axes,
        categories=categories,
        chart_xml=chart_xml_asset,
        workbook=workbook_ref,
        external_workbook_uri=external_uri,
        data_source_status=data_status,
        display_blanks_as=display_blanks or "gap",
        part_graph_root=chart_part,
    )


def _axis_list_with_ctx(chart_space: etree._Element, ctx: ImportContext) -> list[AxisIR]:
    global ctx_hash_placeholder
    ctx_hash_placeholder = ctx.source_hash
    try:
        return _axis_list(chart_space)
    finally:
        ctx_hash_placeholder = ""


def _verify_workbook(
    ctx: ImportContext,
    workbook_ref: Any,
    series: list[ChartSeriesIR],
    chart_part: str,
) -> None:
    """Compare cached chart values with the embedded workbook (spec §10.6 step 6)."""
    try:
        import io

        from openpyxl import load_workbook
    except ImportError:
        return
    try:
        data = ctx.asset_store.read_bytes(workbook_ref)
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - unreadable workbook is a diagnostic
        ctx.add_issue(
            "CHART_WORKBOOK_UNREADABLE",
            f"embedded workbook could not be parsed: {exc}",
            severity=Severity.WARNING,
        )
        return
    try:
        mismatches = 0
        checked = 0
        for ser in series:
            for formula in ser.formula_refs:
                m = _REF_FORMULA_RE.match(formula.strip())
                if not m or "sheet" not in m.groupdict() or m.groupdict()["sheet"] is None:
                    continue
                sheet = m.group("sheet")
                if sheet not in wb.sheetnames:
                    continue
                ws = wb[sheet]
                try:
                    cells = ws[m.group("range").replace("$", "")]
                except Exception:  # noqa: BLE001
                    continue
                cached = [
                    p.value.decimal
                    for p in ser.points
                    if p.value and p.value.decimal is not None
                ]
                flat: list[Any] = []
                rows: list[Any] = list(cells) if isinstance(cells, tuple) else [cells]
                for row in rows:
                    if isinstance(row, tuple):
                        flat.extend(c.value for c in row)
                    else:
                        flat.append(getattr(row, "value", None))
                flat_vals = [str(v) for v in flat if v is not None]
                for cval, wval in zip(cached, flat_vals, strict=False):
                    checked += 1
                    try:
                        from decimal import Decimal

                        if Decimal(cval) != Decimal(wval):
                            mismatches += 1
                    except Exception:  # noqa: BLE001
                        if cval != wval:
                            mismatches += 1
        if mismatches:
            ctx.add_issue(
                "CHART_CACHE_WORKBOOK_CONFLICT",
                f"{mismatches} cached chart values differ from the embedded workbook "
                f"({checked} compared) in {chart_part}",
                severity=Severity.ERROR,
            )
    finally:
        import contextlib

        with contextlib.suppress(Exception):
            wb.close()


# ---------------------------------------------------------------------------
# Other readers (§10.7)
# ---------------------------------------------------------------------------

def read_image(element: etree._Element, ctx: ImportContext, shape_key: object) -> ImagePayload:
    blip = element.find(".//a:blip", namespaces=NS)
    rid = blip.get(qn("r:embed")) or blip.get(qn("r:link")) or "" if blip is not None else ""
    part, external = ctx.rel_target(ctx.slide_part, rid) if rid else (None, None)
    asset = None
    if part and part in ctx.graph.parts_by_name:
        part_ir = ctx.graph.parts_by_name[part]
        if part_ir.asset is not None:
            asset = part_ir.asset
    if asset is None and external:
        ctx.add_issue(
            "IMAGE_LINKED_EXTERNAL",
            f"image relationship is external: {external}",
            severity=Severity.WARNING,
        )
        # Preserve bytes only when part exists; otherwise raise through capability.
    if asset is None:
        raise ValueError(f"image asset not resolvable (rId={rid})")

    src_rect = element.find(".//a:srcRect", namespaces=NS)
    crop = (
        int(src_rect.get("l") or 0) / 100000.0,
        int(src_rect.get("t") or 0) / 100000.0,
        int(src_rect.get("r") or 0) / 100000.0,
        int(src_rect.get("b") or 0) / 100000.0,
    ) if src_rect is not None else (0.0, 0.0, 0.0, 0.0)

    xfrm = element.find(".//a:xfrm", namespaces=NS)
    rotation = float(xfrm.get("rot") or 0) / 60000.0 if xfrm is not None else 0.0
    cnv = element.find(".//p:cNvPr", namespaces=NS)
    alt = cnv.get("descr") if cnv is not None else None
    return ImagePayload(
        asset=asset,
        crop=crop,
        rotation_deg=rotation,
        source_role="unknown",
        alt_text=alt,
        linked_uri=external,
    )


def read_shape(element: etree._Element, ctx: ImportContext, shape_key: object) -> ShapePayload:
    sp_pr = element.find("p:spPr", namespaces=NS)
    geometry_kind: Literal["preset", "freeform"] = "preset"
    geometry_token = None
    path_commands: list[PathCommand] = []
    if sp_pr is not None:
        prst = sp_pr.find("a:prstGeom", namespaces=NS)
        cust = sp_pr.find("a:custGeom", namespaces=NS)
        if prst is not None:
            geometry_kind = "preset"
            geometry_token = prst.get("prst")
        elif cust is not None:
            geometry_kind = "freeform"
            geometry_token = "custom"
            for path in cust.findall(".//a:path", namespaces=NS):
                for cmd in path:
                    local = etree.QName(cmd).localname
                    pts: list[PointPt] = []
                    if local in ("moveTo", "lnTo"):
                        p_node = cmd.find("a:pt", namespaces=NS)
                        if p_node is not None:
                            pts = [
                                PointPt(
                                    x=emu_to_pt(_emu_int(p_node.get("x"))),
                                    y=emu_to_pt(_emu_int(p_node.get("y"))),
                                )
                            ]
                            path_commands.append(
                                PathCommand(op="move" if local == "moveTo" else "line", points=pts)
                            )
                    elif local == "cubicBezTo":
                        for p_node in cmd.findall("a:pt", namespaces=NS):
                            pts.append(
                                PointPt(
                                    x=emu_to_pt(_emu_int(p_node.get("x"))),
                                    y=emu_to_pt(_emu_int(p_node.get("y"))),
                                )
                            )
                        if len(pts) == 3:
                            path_commands.append(PathCommand(op="cubic", points=pts))
                    elif local == "quadBezTo":
                        for p_node in cmd.findall("a:pt", namespaces=NS):
                            pts.append(
                                PointPt(
                                    x=emu_to_pt(_emu_int(p_node.get("x"))),
                                    y=emu_to_pt(_emu_int(p_node.get("y"))),
                                )
                            )
                        if len(pts) == 2:
                            path_commands.append(PathCommand(op="quad", points=pts))
                    elif local == "close":
                        path_commands.append(PathCommand(op="close", points=[]))

    line = LineStyleIR()
    fill_expr = None
    if sp_pr is not None:
        ln = sp_pr.find("a:ln", namespaces=NS)
        if ln is not None:
            line = LineStyleIR(
                color=_fill_color(ln),
                width_emu=_emu_int(ln.get("w")) or None,
                dash=ln.find("a:prstDash", namespaces=NS).get("val")
                if ln.find("a:prstDash", namespaces=NS) is not None else None,
                begin_arrow=ln.find("a:headEnd", namespaces=NS).get("type")
                if ln.find("a:headEnd", namespaces=NS) is not None else None,
                end_arrow=ln.find("a:tailEnd", namespaces=NS).get("type")
                if ln.find("a:tailEnd", namespaces=NS) is not None else None,
            )
        fill_expr = _fill_color(sp_pr)

    text = None
    tx = element.find("p:txBody", namespaces=NS)
    if tx is not None:
        text = read_text_body(tx, ctx, "text", shape_key)
    return ShapePayload(
        geometry_kind=geometry_kind,
        geometry_token=geometry_token,
        path_commands=path_commands,
        text=text,
        source_line=line,
        source_fill=fill_expr,
    )


def read_connector(
    element: etree._Element, ctx: ImportContext, shape_key: object
) -> ConnectorPayload:
    cnv_cxn = element.find(".//p:cNvCxnSpPr", namespaces=NS)
    from_id = to_id = None
    from_site = to_site = None
    evidence: list[str] = []
    if cnv_cxn is not None:
        st = cnv_cxn.find("a:stCxn", namespaces=NS)
        en = cnv_cxn.find("a:endCxn", namespaces=NS)
        if st is not None:
            from_id = st.get("id")
            from_site = int(st.get("idx") or 0)
            evidence.append(f"stCxn id={from_id} idx={from_site}")
        if en is not None:
            to_id = en.get("id")
            to_site = int(en.get("idx") or 0)
            evidence.append(f"endCxn id={to_id} idx={to_site}")

    sp_pr = element.find("p:spPr", namespaces=NS)
    arrow_start = arrow_end = "none"
    if sp_pr is not None:
        ln = sp_pr.find("a:ln", namespaces=NS)
        if ln is not None:
            head = ln.find("a:headEnd", namespaces=NS)
            tail = ln.find("a:tailEnd", namespaces=NS)
            arrow_start = head.get("type") if head is not None else "none"
            arrow_end = tail.get("type") if tail is not None else "none"

    text = None
    tx = element.find("p:txBody", namespaces=NS)
    if tx is not None:
        text = read_text_body(tx, ctx, "label", shape_key)

    return ConnectorPayload(
        from_object_id=None if from_id is None else f"shapeid:{from_id}",
        to_object_id=None if to_id is None else f"shapeid:{to_id}",
        from_site=from_site,
        to_site=to_site,
        arrow_start=arrow_start,
        arrow_end=arrow_end,
        relation_kind="unknown",
        connection_evidence=evidence,
        text=text,
    )


def read_unknown(element: etree._Element, ctx: ImportContext, shape_key: object) -> UnknownPayload:
    xml_blob = etree.tostring(element, encoding="utf-8")
    asset = ctx.asset_store.put(xml_blob, "application/xml")
    object_type = etree.QName(element).localname
    texts: list[TextPayload] = []
    found = element.findall(".//p:txBody", namespaces=NS)
    found += element.findall(".//a:txBody", namespaces=NS)
    for tx in found:
        try:
            texts.append(read_text_body(tx, ctx, "unknown/text", shape_key))
        except Exception:  # noqa: BLE001 - unknown structures may not parse
            continue
    graphic = element.find(".//a:graphicData", namespaces=NS)
    uri = graphic.get("uri") if graphic is not None else None
    return UnknownPayload(
        raw_shape_xml=asset,
        part_graph_root=ctx.slide_part,
        object_type=f"{object_type}:{uri}" if uri else object_type,
        extracted_text=texts,
        capability=Capability(
            level="unsupported",
            reason="graphicData type is not part of the supported reader set",
            missing_features=[uri or object_type],
        ),
    )


# ---------------------------------------------------------------------------
# Shape tree walk (§10.3)
# ---------------------------------------------------------------------------

_SHAPE_TAGS = {
    "sp": "shape",
    "pic": "image",
    "cxnSp": "connector",
    "grpSp": "group",
    "graphicFrame": "graphicFrame",
    "contentPart": "unknown",
}


def _register_shape_id(ctx: ImportContext, shape_id: int, element: etree._Element) -> bool:
    seen = ctx.shape_ids_by_slide.setdefault(ctx.slide_part, set())
    if shape_id in seen:
        ctx.add_issue(
            "SHAPE_ID_NOT_UNIQUE",
            f"duplicate cNvPr id {shape_id} in {ctx.slide_part}",
            severity=Severity.ERROR,
        )
        return False
    seen.add(shape_id)
    return True


def _cncv(element: etree._Element) -> tuple[int, str, bool]:
    cnv = element.find(".//p:cNvPr", namespaces=NS)
    if cnv is None:
        return 0, "", True
    shape_id = int(cnv.get("id") or 0)
    name = cnv.get("name") or ""
    hidden = cnv.get("hidden") == "1"
    return shape_id, name, hidden


def _local_box(element: etree._Element) -> tuple[RectEMU | None, dict[str, Any] | None]:
    xfrm = element.find("./p:spPr/a:xfrm", namespaces=NS)
    if xfrm is None:
        xfrm = element.find("./p:xfrm", namespaces=NS)
    if xfrm is None:
        xfrm = element.find(".//a:xfrm", namespaces=NS)
    box = _xfrm(xfrm)
    if box is None:
        return None, None
    return RectEMU(x=box["x"], y=box["y"], w=max(box["cx"], 0), h=max(box["cy"], 0)), box


def walk_shape_tree(
    element: etree._Element,
    slide_ref: str,
    parent_transform: Affine2D,
    group_path: tuple[int, ...],
    ctx: ImportContext,
) -> list[ObjectIR]:
    """Walk one spTree level and return typed ObjectIRs (spec §10.3)."""
    objects: list[ObjectIR] = []
    children = [c for c in element if isinstance(c.tag, str)]
    z = 0
    for child in children:
        local = etree.QName(child).localname
        if local not in _SHAPE_TAGS:
            continue
        z_order = z
        z += 1
        shape_id, name, hidden = _cncv(child)
        if shape_id and not _register_shape_id(ctx, shape_id, child):
            continue
        key = (shape_id, list(group_path))
        local_box, box = _local_box(child)
        local_rot = _rotation_flip_matrix(box) if box else Affine2D()
        slide_transform = compose_affine(parent_transform, local_rot)
        source_ref = SourceRef(
            deck_sha256=ctx.source_hash,
            slide_part=slide_ref,
            shape_id=shape_id,
            group_path=list(group_path),
            subpath=f"sp/{shape_id}",
            xml_sha256=sha256_bytes(etree.tostring(child)),
        )

        if local == "grpSp":
            if box is None:
                ctx.add_issue(
                    "GROUP_WITHOUT_XFRM", f"group {shape_id} has no xfrm",
                    object_ids=[str(shape_id)],
                )
                continue
            child_matrix = compose_affine(parent_transform, group_child_matrix(box))
            child_path = (*group_path, shape_id)
            grand_children = walk_shape_tree(child, slide_ref, child_matrix, child_path, ctx)
            tx = child.find("p:txBody", namespaces=NS)
            text = read_text_body(tx, ctx, "text", key) if tx is not None else None
            payload = GroupPayload(
                children=grand_children,
                local_transform=group_child_matrix(box),
                source_child_box=RectEMU(
                    x=box.get("chX", 0), y=box.get("chY", 0),
                    w=max(box.get("chCX", 0), 0), h=max(box.get("chCY", 0), 0),
                ),
                text=text,
            )
            objects.append(
                ObjectIR(
                    id=stable_id(ctx.source_hash, slide_ref, list(group_path), shape_id, "group"),
                    source_ref=source_ref,
                    kind="group",
                    payload=payload,
                    local_box=local_box,
                    slide_transform=slide_transform,
                    z_order=z_order,
                    visible=not hidden,
                    name=name,
                    capability=Capability(level="native_full"),
                    semantic_significance="unknown",
                )
            )
            continue

        if local == "graphicFrame":
            obj = _read_graphic_frame(
                child, ctx, key, source_ref, local_box, slide_transform, z_order, name
            )
            if obj is not None:
                objects.append(obj)
            continue

        try:
            if local == "sp":
                has_tx_body = child.find("p:txBody", namespaces=NS) is not None
                cnv_sp = child.find(".//p:cNvSpPr", namespaces=NS)
                is_textbox = cnv_sp is not None and cnv_sp.get("txBox") == "1"
                sp_pr = child.find("p:spPr", namespaces=NS)
                has_geom = sp_pr is not None and (
                    sp_pr.find("a:prstGeom", namespaces=NS) is not None
                    or sp_pr.find("a:custGeom", namespaces=NS) is not None
                )
                payload: Any
                if has_tx_body and (is_textbox or not has_geom):
                    tx_body = child.find("p:txBody", namespaces=NS)
                    assert tx_body is not None
                    payload = read_text_body(tx_body, ctx, "text", key)
                    kind = "text"
                else:
                    payload = read_shape(child, ctx, key)
                    kind = "shape"
            elif local == "pic":
                payload = read_image(child, ctx, key)
                kind = "image"
            elif local == "cxnSp":
                payload = read_connector(child, ctx, key)
                kind = "connector"
            else:
                payload = read_unknown(child, ctx, key)
                kind = "unknown"
        except Exception as exc:  # noqa: BLE001 - convert to explicit unsupported object
            ctx.add_issue(
                "READER_FAILED",
                f"{local} #{shape_id}: {exc}",
                severity=Severity.ERROR,
                object_ids=[str(shape_id)],
            )
            payload = read_unknown(child, ctx, key)
            kind = "unknown"

        objects.append(
            ObjectIR(
                id=stable_id(ctx.source_hash, slide_ref, list(group_path), shape_id, local),
                source_ref=source_ref,
                kind=kind,  # type: ignore[arg-type]
                payload=payload,
                local_box=local_box,
                slide_transform=slide_transform,
                z_order=z_order,
                visible=not hidden,
                name=name,
                capability=Capability(level="native_full" if kind != "unknown" else "unsupported"),
                semantic_significance="unknown",
            )
        )
    return objects


def _read_graphic_frame(
    frame: etree._Element,
    ctx: ImportContext,
    key: tuple[int, list[int]],
    source_ref: SourceRef,
    local_box: RectEMU | None,
    slide_transform: Affine2D,
    z_order: int,
    name: str,
) -> ObjectIR | None:
    graphic = frame.find(".//a:graphicData", namespaces=NS)
    uri = graphic.get("uri") or "" if graphic is not None else ""
    shape_id = key[0]
    payload: Any
    kind: str
    capability = Capability(level="native_full")
    if graphic is not None and graphic.find("a:tbl", namespaces=NS) is not None:
        payload = read_table(frame, ctx, key)
        kind = "table"
    elif graphic is not None and graphic.find("c:chart", namespaces=NS) is not None:
        payload = read_chart(frame, ctx, key)
        kind = "chart"
        if isinstance(payload, ChartPayload) and payload.data_source_status == "missing":
            capability = Capability(
                level="native_partial", reason="chart data source missing",
                missing_features=["workbook", "cache"],
            )
    elif uri.endswith(("/diagram", "/chart")):
        payload = read_unknown(frame, ctx, key)
        kind = "unknown"
        capability = payload.capability
    else:
        payload = read_unknown(frame, ctx, key)
        kind = "unknown"
        capability = payload.capability

    return ObjectIR(
        id=stable_id(ctx.source_hash, ctx.slide_part, list(key[1]), shape_id, kind),
        source_ref=source_ref,
        kind=kind,  # type: ignore[arg-type]
        payload=payload,
        local_box=local_box,
        slide_transform=slide_transform,
        z_order=z_order,
        visible=True,
        name=name,
        capability=capability,
        semantic_significance="unknown",
    )


def extract_notes(slide_part: str, ctx: ImportContext) -> list[ParagraphIR]:
    """Read speaker notes paragraphs from the notes slide part (§10.7)."""
    for rel in ctx.graph.outgoing.get(slide_part, []):
        if rel.rel_type.endswith("/notesSlide") and rel.resolved_part:
            notes_part = rel.resolved_part
            asset = ctx.graph.parts_by_name.get(notes_part)
            if asset is None or asset.asset is None:
                return []
            xml = ctx.asset_store.read_bytes(asset.asset)
            root = etree.fromstring(
                xml, parser=etree.XMLParser(resolve_entities=False, no_network=True)
            )
            paragraphs: list[ParagraphIR] = []
            for idx, sp in enumerate(root.findall(".//p:sp", namespaces=NS)):
                ph = sp.find(".//p:ph", namespaces=NS)
                if ph is not None and ph.get("type") not in (None, "body"):
                    continue
                tx = sp.find("p:txBody", namespaces=NS)
                if tx is None:
                    continue
                payload = read_text_body(tx, ctx, f"notes/sp{idx}", 0)
                paragraphs.extend(payload.paragraphs)
            return paragraphs
    return []


def is_hidden_slide(slide_root: etree._Element) -> bool:
    return slide_root.get("show") == "0"
