"""Stage 7 §9: second input mode — PDF/DOC/DOCX/MD/TXT/prompt create.

PPTX keeps using the existing importer. Document inputs become an ordered
``SourceDocument`` (blocks, tables, images, provenance), then a
``narrative_plan`` model task, then the SAME SemanticDeck -> ScenePlan ->
measure -> native emit -> checks -> render/review/selection/acceptance
pipeline as restyle. No HTML-to-image fallback, no fabricated figures.

Errors: INPUT_KIND_UNSUPPORTED, CONVERSION_TOOL_MISSING, CONVERSION_FAILED,
DOCUMENT_EXTRACTION_INCOMPLETE, OCR_UNCERTAIN_REQUIRED_FACT,
SOURCE_REFERENCE_MISSING, NARRATIVE_REQUIRED_FACT_MISSING,
DOCUMENT_DATA_AMBIGUOUS.
"""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
from pathlib import Path

from pydantic import BaseModel, Field

from .models import (
    InputKind,
    InputSpec,
    ModelTask,
    ModelTaskKind,
    ModelTaskResponse,
    StrictModel,
)


class IngestError(ValueError):
    """Typed ingest failure with a stable error code prefix."""


class SourceBlock(BaseModel):
    model_config = StrictModel

    block_id: str
    kind: str = "paragraph"
    text: str = ""
    page_index: int | None = None
    provenance: str = ""


class SourceAsset(BaseModel):
    model_config = StrictModel

    asset_id: str
    kind: str = "image"
    path: str = ""
    sha256: str = ""
    page_index: int | None = None


class SourceDocument(BaseModel):
    """Ordered, provenance-carrying document extraction (§9.1)."""

    model_config = StrictModel

    schema_version: str = "source-document/v1"
    input_kind: InputKind
    source_path: str = ""
    content_sha256: str = ""
    source_sha256: str = ""
    blocks: list[SourceBlock] = Field(default_factory=list)
    tables: list[dict] = Field(default_factory=list)
    assets: list[SourceAsset] = Field(default_factory=list)
    extraction_issues: list[str] = Field(default_factory=list)
    conversion: dict = Field(default_factory=dict)


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def detect_kind(path: Path) -> InputKind:
    suffix = path.suffix.lower()
    mapping = {".pptx": InputKind.PPTX, ".ppt": InputKind.PPT,
               ".pdf": InputKind.PDF, ".docx": InputKind.DOCX,
               ".doc": InputKind.DOC, ".md": InputKind.MD,
               ".txt": InputKind.TXT}
    kind = mapping.get(suffix)
    if kind is None:
        raise IngestError(f"INPUT_KIND_UNSUPPORTED: {path.suffix}")
    return kind


def ingest_input(spec: InputSpec, run_dir: Path) -> SourceDocument:
    """§9.1 entry: route by kind, write inputs/source_manifest.json."""
    run_dir = Path(run_dir)
    (run_dir / "inputs").mkdir(parents=True, exist_ok=True)
    if spec.kind == InputKind.PROMPT:
        if not (spec.prompt_text or "").strip():
            raise IngestError("INPUT_KIND_UNSUPPORTED: empty prompt")
        doc = extract_prompt(spec.prompt_text or "")
    elif spec.kind in (InputKind.PPTX, InputKind.PPT):
        if spec.path is None:
            raise IngestError("SOURCE_REFERENCE_MISSING: path required")
        doc = extract_pptx(_resolve(spec, spec.path), run_dir)
    else:
        if spec.path is None:
            raise IngestError("SOURCE_REFERENCE_MISSING: path required")
        doc = extract_document(_resolve(spec, spec.path), spec.kind, run_dir)
    (run_dir / "inputs" / "source_manifest.json").write_text(
        doc.model_dump_json(indent=1), encoding="utf-8")
    return doc


def _resolve(spec: InputSpec, path: Path) -> Path:
    p = Path(path)
    if not p.is_absolute():
        p = Path(spec.project_root) / p
    if not p.is_file():
        raise IngestError(f"SOURCE_REFERENCE_MISSING: {p}")
    return p


# ---------------------------------------------------------------------------
# PPTX (+ legacy PPT via local conversion)
# ---------------------------------------------------------------------------

def extract_pptx(path: Path, run_dir: Path) -> SourceDocument:
    data = path.read_bytes()
    if path.suffix.lower() == ".ppt":
        converted, tool = _convert_office(path, run_dir, "pptx")
        return SourceDocument(
            input_kind=InputKind.PPT, source_path=str(path),
            content_sha256=_sha(converted.read_bytes()),
            source_sha256=_sha(data),
            blocks=[SourceBlock(
                block_id="ppt-ledger",
                kind="presentation",
                text=f"converted legacy PPT via {tool}",
                provenance=str(converted))],
            conversion={"tool": tool,
                        "input_sha256": _sha(data),
                        "output_sha256": _sha(converted.read_bytes()),
                        "uncertainty": "layout fidelity of legacy "
                                       "conversion unverified"})
    # PPTX: the existing importer owns IR/ledger; ingest records identity.
    return SourceDocument(
        input_kind=InputKind.PPTX, source_path=str(path),
        content_sha256=_sha(data), source_sha256=_sha(data),
        blocks=[SourceBlock(block_id="pptx-ledger", kind="presentation",
                            text="source IR via gpn.importer",
                            provenance=str(path))])


def _convert_office(path: Path, run_dir: Path, fmt: str) -> tuple[Path, str]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise IngestError("CONVERSION_TOOL_MISSING: install LibreOffice "
                          "for legacy DOC/PPT conversion")
    outdir = run_dir / "inputs" / "converted"
    outdir.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run([soffice, "--headless", "--convert-to", fmt,
                        "--outdir", str(outdir), str(path)],
                       check=True, capture_output=True, timeout=300)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        raise IngestError(f"CONVERSION_FAILED: {exc}") from exc
    outputs = sorted(outdir.glob(f"{path.stem}.{fmt}"))
    if not outputs:
        raise IngestError(f"CONVERSION_FAILED: no output for {path.name}")
    version = subprocess.run([soffice, "--version"], capture_output=True,
                             text=True, timeout=60)
    return outputs[0], version.stdout.strip() or soffice


# ---------------------------------------------------------------------------
# documents
# ---------------------------------------------------------------------------

def extract_document(path: Path, kind: InputKind,
                     run_dir: Path) -> SourceDocument:
    data = path.read_bytes()
    if kind == InputKind.PDF:
        return _extract_pdf(path, data, run_dir)
    if kind in (InputKind.DOCX, InputKind.DOC):
        return _extract_word(path, kind, data, run_dir)
    if kind in (InputKind.MD, InputKind.TXT):
        return _extract_text(path, kind, data)
    raise IngestError(f"INPUT_KIND_UNSUPPORTED: {kind}")


def extract_prompt(text: str) -> SourceDocument:
    data = text.encode("utf-8")
    blocks = [SourceBlock(block_id=f"prompt-p{i}", kind="paragraph",
                          text=line.strip(), provenance="user prompt")
              for i, line in enumerate(text.split("\n")) if line.strip()]
    if not blocks:
        raise IngestError("INPUT_KIND_UNSUPPORTED: empty prompt")
    return SourceDocument(
        input_kind=InputKind.PROMPT, content_sha256=_sha(data),
        source_sha256=_sha(data), blocks=blocks)


def _extract_text(path: Path, kind: InputKind, data: bytes) -> SourceDocument:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise IngestError(
            f"DOCUMENT_EXTRACTION_INCOMPLETE: {path.name} not UTF-8") from exc
    blocks: list[SourceBlock] = []
    current_section = ""
    for i, line in enumerate(text.split("\n")):
        stripped = line.strip()
        if not stripped:
            continue
        if kind == InputKind.MD and stripped.startswith("#"):
            current_section = stripped.lstrip("#").strip()
            blocks.append(SourceBlock(
                block_id=f"{path.stem}-h{i}", kind="heading",
                text=current_section, provenance=f"{path.name}:{i + 1}"))
        elif stripped.startswith(("- ", "* ", "1. ", "2. ", "3. ")):
            blocks.append(SourceBlock(
                block_id=f"{path.stem}-li{i}", kind="list_item",
                text=stripped, provenance=f"{path.name}:{i + 1}"))
        else:
            blocks.append(SourceBlock(
                block_id=f"{path.stem}-p{i}", kind="paragraph",
                text=stripped, provenance=f"{path.name}:{i + 1}"))
    return SourceDocument(input_kind=kind, source_path=str(path),
                          content_sha256=_sha(data),
                          source_sha256=_sha(data), blocks=blocks)


def _extract_pdf(path: Path, data: bytes,
                 run_dir: Path) -> SourceDocument:
    try:
        from pypdf import PdfReader  # noqa: PLC0415
    except ImportError as exc:
        raise IngestError(
            "DOCUMENT_EXTRACTION_INCOMPLETE: pypdf not installed") from exc
    reader = PdfReader(str(path))
    blocks: list[SourceBlock] = []
    issues: list[str] = []
    for page_index, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            issues.append(f"page {page_index}: {exc}")
            continue
        if not text.strip():
            issues.append(f"page {page_index}: no extractable text "
                          f"(scanned? OCR via harness vision task)")
            continue
        for j, para in enumerate(text.split("\n")):
            if para.strip():
                blocks.append(SourceBlock(
                    block_id=f"pdf-p{page_index}-{j}", kind="paragraph",
                    text=para.strip(), page_index=page_index,
                    provenance=f"{path.name}#page={page_index + 1}"))
    if not blocks:
        raise IngestError(
            "DOCUMENT_EXTRACTION_INCOMPLETE: no text blocks extracted; "
            "scanned PDF needs OCR task before narrative planning")
    return SourceDocument(input_kind=InputKind.PDF, source_path=str(path),
                          content_sha256=_sha(data),
                          source_sha256=_sha(data), blocks=blocks,
                          extraction_issues=issues)


def _extract_word(path: Path, kind: InputKind, data: bytes,
                  run_dir: Path) -> SourceDocument:
    working = path
    conversion: dict = {}
    if kind == InputKind.DOC:
        working, tool = _convert_office(path, run_dir, "docx")
        conversion = {"tool": tool, "input_sha256": _sha(data),
                      "output_sha256": _sha(working.read_bytes())}
    try:
        from docx import Document  # noqa: PLC0415
    except ImportError as exc:
        raise IngestError(
            "DOCUMENT_EXTRACTION_INCOMPLETE: python-docx not "
            "installed") from exc
    document = Document(str(working))
    blocks: list[SourceBlock] = []
    tables: list[dict] = []
    for i, para in enumerate(document.paragraphs):
        if not para.text.strip():
            continue
        kind_name = "heading" if para.style.name.startswith("Heading") \
            else "paragraph"
        blocks.append(SourceBlock(
            block_id=f"docx-p{i}", kind=kind_name, text=para.text.strip(),
            provenance=f"{path.name}#p{i}"))
    for ti, table in enumerate(document.tables):
        rows = [[cell.text.strip() for cell in row.cells]
                for row in table.rows]
        tables.append({"table_id": f"docx-t{ti}", "rows": rows,
                       "provenance": f"{path.name}#table{ti}"})
        blocks.append(SourceBlock(
            block_id=f"docx-t{ti}", kind="table",
            text=" | ".join(" / ".join(r) for r in rows),
            provenance=f"{path.name}#table{ti}"))
    for rel in document.part.rels.values():
        if "image" in rel.target_ref:
            target = rel.target_part
            digest = _sha(target.blob)
            asset_path = run_dir / "inputs" / "assets" / \
                f"{digest[:16]}-{Path(target.partname).name}"
            asset_path.parent.mkdir(parents=True, exist_ok=True)
            asset_path.write_bytes(target.blob)
            blocks.append(SourceBlock(
                block_id=f"docx-img-{digest[:8]}", kind="image",
                text=f"[image {digest[:12]}]",
                provenance=f"{path.name}#{target.partname}"))
    if not blocks:
        raise IngestError("DOCUMENT_EXTRACTION_INCOMPLETE: empty document")
    return SourceDocument(input_kind=kind, source_path=str(path),
                          content_sha256=_sha(data),
                          source_sha256=_sha(data), blocks=blocks,
                          tables=tables, conversion=conversion)


# ---------------------------------------------------------------------------
# semantic task (§8.3: request_semantic_projection)
# ---------------------------------------------------------------------------

def source_object_inventory(run_dir: Path) -> list[dict]:
    """Compact source object view for semantic classification (§6.2).

    The model classifies *objects* (not sentences): every text-bearing
    object becomes a required unit, every arrow-bearing shape a required
    relation, every other shape needs an explicit decision.
    """
    ir_path = run_dir / "source_ir.json"
    if not ir_path.is_file():
        return []
    ir = json.loads(ir_path.read_text(encoding="utf-8"))
    out: list[dict] = []
    for slide in ir.get("slides", []):
        objects: list[dict] = []
        for obj in slide.get("objects", []):
            payload = obj.get("payload") or {}
            text_obj = payload.get("text") or {}
            paras = payload.get("paragraphs")
            if paras is None:
                paras = text_obj.get("paragraphs", []) or []
            text = "".join(
                (run.get("text") or "") for para in paras
                for run in (para.get("runs") or []))
            line = payload.get("source_line") or {}
            objects.append({
                "object_id": str(obj.get("id", "")),
                "kind": str(obj.get("kind", "")),
                "name": str(obj.get("name", "")),
                "box": obj.get("local_box"),
                "geometry_token": payload.get("geometry_token"),
                "begin_arrow": line.get("begin_arrow"),
                "end_arrow": line.get("end_arrow"),
                "text": text,
                "asset_sha256": (payload.get("asset") or {}).get("sha256"),
            })
        out.append({"source_slide_id": str(slide.get("id", "")),
                    "hidden": bool(slide.get("hidden")),
                    "notes_text": " ".join(
                        "".join((r.get("text") or "")
                                for r in (n.get("runs") or []))
                        for n in (slide.get("notes") or [])),
                    "objects": objects})
    return out


def semantic_task_for(spec: InputSpec, run_dir: Path,
                      input_sha: str) -> ModelTask:
    """Build the semantic_analysis / narrative_plan task packet (§9.5)."""
    manifest_path = run_dir / "inputs" / "source_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    is_create = spec.kind != InputKind.PPTX and spec.kind != InputKind.PPT
    kind = ModelTaskKind.NARRATIVE_PLAN if is_create \
        else ModelTaskKind.SEMANTIC_ANALYSIS
    schema = "narrative-plan/v1" if is_create else "semantic-deck/v1"
    packet = {
        "stage": kind.value,
        "input_kind": spec.kind.value,
        "source_manifest": manifest,
        "requested_slide_count": spec.requested_slide_count,
        "source_object_inventory": source_object_inventory(run_dir),
        "response_schema": schema,
        "required_fields": ["slides", "units", "relations"]
        if not is_create else
        ["narrative", "slides", "required_fact_coverage"],
        "transformation_decisions_rule": (
            "one decision per source object: preserve (bound to a target "
            "node), recompose/replace (with target_relation_ids when the "
            "source object carried an arrow), or omit (only evidence-backed "
            "nonsemantic decoration, with a reason). Absence of text is "
            "never a reason; arrow-bearing shapes are relations."),
        "rules": ["no invented figures/conclusions/savings",
                  "numbers/units/caveats preserved with provenance",
                  "create: requested count honored or infeasible proof"],
    }
    rel = f"model_requests/{kind.value}.json"
    (run_dir / rel).parent.mkdir(parents=True, exist_ok=True)
    (run_dir / rel).write_text(json.dumps(packet, ensure_ascii=False,
                                          indent=1), encoding="utf-8")
    return ModelTask(
        request_id=f"T-semantic-{input_sha[:12]}",
        kind=kind,
        input_fingerprint=input_sha,
        response_schema_id=schema,
        packet=rel,
        expected_response_units=packet["required_fields"])


def validate_semantic_response(response: ModelTaskResponse,
                               source: SourceDocument) -> list[str]:
    """§9.6: narrative must cover source facts, invent nothing checkable."""
    errors: list[str] = []
    payload = response.payload
    if not payload.get("slides"):
        errors.append("NARRATIVE_REQUIRED_FACT_MISSING: no slides planned")
        return errors
    source_text = " ".join(b.text for b in source.blocks)
    import re
    numbers = set(re.findall(r"\d[\d\s.,]*", source_text))
    planned = json.dumps(payload, ensure_ascii=False)
    missing = [n.strip() for n in numbers
               if n.strip() and len(n.strip()) >= 2
               and n.strip() not in planned]
    if missing:
        errors.append("NARRATIVE_REQUIRED_FACT_MISSING: source literals "
                      f"absent from plan: {missing[:5]}")
    return errors


__all__ = ["IngestError", "SourceBlock", "SourceAsset", "SourceDocument",
           "detect_kind", "ingest_input", "extract_pptx", "extract_document",
           "extract_prompt", "semantic_task_for", "source_object_inventory",
           "validate_semantic_response"]
