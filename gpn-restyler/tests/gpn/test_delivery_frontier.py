"""Causal orchestration regressions. Native/Chromium boundaries are explicit."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import zipfile
from pathlib import Path

import pytest

from slides_cli.gpn import restyle_runner as runner
from slides_cli.gpn.design import build_slide_design_brief
from slides_cli.gpn.models import (
    CandidateAttemptResult,
    CompiledOntology,
    InputSpec,
    ModelTaskKind,
    ModelTaskResponse,
    RoleStyle,
    SemanticDeck,
    SemanticSlide,
    SemanticUnit,
    SourceEvidenceRef,
)
from slides_cli.gpn.scene_jobs import comparison_coverage
from slides_cli.gpn.source_images import source_slide_images

TOOL = Path(__file__).resolve().parents[2] / "tools/harness_batch.py"
spec = importlib.util.spec_from_file_location("frontier_dispatch", TOOL)
dispatch = importlib.util.module_from_spec(spec)
spec.loader.exec_module(dispatch)


def write(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False))


def case(root, count=3):
    deck = SemanticDeck(
        slides=[
            SemanticSlide(
                source_slide_id=f"s{i}",
                units=[
                    SemanticUnit(
                        id=f"u{i}",
                        kind="text",
                        source_refs=[
                            SourceEvidenceRef(
                                artifact_id="source",
                                relative_path="source_ir.json",
                                object_id=f"o{i}",
                            )
                        ],
                    )
                ],
            )
            for i in range(count)
        ]
    )
    write(root / "inputs/semantic_deck.json", deck.model_dump(mode="json"))
    write(
        root / "source_ir.json",
        {"slides": [{"id": f"s{i}", "objects": [{"id": f"o{i}"}]} for i in range(count)]},
    )
    rules = CompiledOntology(roles={"body": RoleStyle(role="body", default_size_pt=12)})
    write(root / "compiled_rules.json", rules.model_dump(mode="json"))
    state = runner.RunState(
        run_id="r",
        run_dir=str(root),
        stage="design_briefs",
        input_spec=InputSpec(
            kind="prompt",
            prompt_text="Test-only orchestration",
            project_root=root,
            output_path=root / "out.pptx",
        ),
    )
    for slide in deck.slides:
        state.task_queue.append(
            build_slide_design_brief(
                slide, [], run_dir=root, request_id=f"brief-{slide.source_slide_id}"
            )
        )
    runner.save_state(state)
    return state


def response(task, payload):
    return ModelTaskResponse(
        request_id=task.request_id,
        input_fingerprint=task.input_fingerprint,
        schema_version_id=task.response_schema_id,
        actual_model="test-only authored fixture",
        payload=payload,
    )


def brief(sid="s0", uid="u0"):
    return {
        "slide_id": sid,
        "communication_job": "Explain the claim",
        "intended_takeaway": "The source claim remains exact",
        "required_units": [uid],
        "required_relations": [],
        "semantic_groups": [{"group_id": "g", "role": "primary", "unit_ids": [uid]}],
        "primary_group_id": "g",
        "visual_encoding_candidates": [],
    }


def concepts():
    return {
        "concepts": [
            {
                "concept_id": f"c{i}",
                "visual_argument": f"Test visual argument {i}",
                "group_topology": {"topology": topology},
                "source_to_visual_mapping": [{"source_unit_ids": ["u0"], "visual_element": "g"}],
            }
            for i, topology in enumerate(("typographic", "evidence_focus", "comparison"), 1)
        ]
    }


@pytest.mark.parametrize("count", [3, 41])
def test_first_brief_unlocks_concepts_while_other_briefs_pending(tmp_path, count):
    state = case(tmp_path, count)
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    assert (
        sum(
            t.status == "pending" and t.kind == ModelTaskKind.DESIGN_BRIEF for t in state.task_queue
        )
        == count - 1
    )
    assert sum(t.kind == ModelTaskKind.VISUAL_CONCEPT_GENERATION for t in state.task_queue) == 1
    assert state.tasks_used == 1
    before = [(t.request_id, t.input_fingerprint) for t in state.task_queue]
    runner.advance_run(state, [])
    assert before == [(t.request_id, t.input_fingerprint) for t in state.task_queue]


def test_accepted_concept_set_unlocks_scenes_without_waiting_other_slides(tmp_path):
    state = case(tmp_path)
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    task = next(t for t in state.task_queue if t.kind == ModelTaskKind.VISUAL_CONCEPT_GENERATION)
    runner.advance_run(state, [response(task, concepts())])
    scenes = [t for t in state.task_queue if t.kind == ModelTaskKind.SCENE_GENERATION]
    assert len(scenes) == 3
    assert all(t.candidate_key.source_slide_id == "s0" for t in scenes)
    assert (
        sum(
            t.status == "pending" and t.kind == ModelTaskKind.DESIGN_BRIEF for t in state.task_queue
        )
        == 2
    )
    before = [(t.request_id, t.input_fingerprint) for t in state.task_queue]
    runner.advance_run(state, [])
    assert before == [(t.request_id, t.input_fingerprint) for t in state.task_queue]


def test_late_brief_is_materialized_even_while_another_slide_is_in_review(tmp_path):
    state = case(tmp_path)
    state.stage = "visual_review"
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    assert (tmp_path / "design_briefs/s0.json").is_file()
    assert any(t.kind == ModelTaskKind.VISUAL_CONCEPT_GENERATION for t in state.task_queue)


def test_answered_brief_can_resume_a_blocked_run_without_changing_request(tmp_path):
    state = case(tmp_path)
    task = state.task_queue[0]
    state.stage = "blocked"
    state.terminal_reason = "legacy stage barrier"
    runner.advance_run(state, [response(task, brief())])
    assert task.status == "consumed"
    assert state.stage == "visual_concepts" and state.terminal_reason == ""


def test_one_visual_review_unlocks_its_design_review_while_other_reviews_pending(
    tmp_path, monkeypatch
):
    from slides_cli.gpn.design_frontier import pump_current_reviews
    from slides_cli.gpn.models import ModelTask

    state = case(tmp_path)
    candidate = tmp_path / "scenes/job-a/candidate.pptx"
    candidate.parent.mkdir(parents=True)
    candidate.write_bytes(b"test boundary bytes, not a native slide assertion")
    sha = hashlib.sha256(candidate.read_bytes()).hexdigest()
    visual = ModelTask(request_id="visual-a", kind=ModelTaskKind.VISUAL_REVIEW, status="consumed")
    other = ModelTask(request_id="visual-b", kind=ModelTaskKind.VISUAL_REVIEW)
    state.task_queue.extend([visual, other])
    state.responses.append(
        response(
            visual,
            {
                "deck_key": "job-a",
                "candidate_sha256": sha,
                "verdict": "accept",
                "pages": [],
            },
        )
    )
    state.deck_candidates = [{"deck_key": "job-a"}]
    design = ModelTask(request_id="design-a", kind=ModelTaskKind.DESIGN_REVIEW)
    monkeypatch.setattr(runner, "_design_review_task", lambda state, job: design)
    pump_current_reviews(state)
    assert other.status == "pending"
    assert any(t.request_id == "design-a" for t in state.task_queue)
    saved = dispatch.read(candidate.parent / "visual_review.json")
    assert saved["candidate_sha256"] == sha


def test_old_hash_review_never_unlocks_current_candidate_design_review(tmp_path, monkeypatch):
    from slides_cli.gpn.design_frontier import pump_current_reviews
    from slides_cli.gpn.models import ModelTask

    state = case(tmp_path)
    candidate = tmp_path / "scenes/job-a/candidate.pptx"
    candidate.parent.mkdir(parents=True)
    candidate.write_bytes(b"new-candidate")
    task = ModelTask(request_id="visual-a", kind=ModelTaskKind.VISUAL_REVIEW, status="consumed")
    state.task_queue.append(task)
    state.responses.append(
        response(
            task,
            {
                "deck_key": "job-a",
                "candidate_sha256": "old-hash",
                "verdict": "accept",
                "pages": [],
            },
        )
    )
    state.deck_candidates = [{"deck_key": "job-a"}]

    def unexpected(*args):
        raise AssertionError("stale review must not unlock anything")

    monkeypatch.setattr(runner, "_design_review_task", unexpected)
    assert not pump_current_reviews(state)
    assert not (candidate.parent / "visual_review.json").exists()


def test_first_scene_reaches_builder_with_two_unanswered_briefs(tmp_path, monkeypatch):
    state = case(tmp_path)
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    task = next(t for t in state.task_queue if t.kind == ModelTaskKind.VISUAL_CONCEPT_GENERATION)
    runner.advance_run(state, [response(task, concepts())])
    scene_task = next(t for t in state.task_queue if t.kind == ModelTaskKind.SCENE_GENERATION)
    seen = []

    def native_boundary(state, key, selection, config, **scope):
        # This test proves dispatch to the real isolated job builder, NOT
        # Chromium measurement or native emission. Those are separate gates.
        seen.append(selection)
        return CandidateAttemptResult(eligible=False)

    monkeypatch.setattr(runner, "_build_deck_candidate", native_boundary)
    monkeypatch.setattr(runner, "_config_from_state", lambda state: None)
    plan = {
        "slide_id": "s0",
        "candidate_id": "c1",
        "content_refs": ["u0"],
        "reading_order": ["z"],
        "dominant_ref": "z",
        "zones": [
            {
                "zone_id": "z",
                "role": "body",
                "content_refs": ["u0"],
                "x_pct": 0,
                "y_pct": 0,
                "w_pct": 100,
                "h_pct": 100,
            }
        ],
    }
    runner.advance_run(state, [response(scene_task, {"scene_plan": plan})])
    assert seen == [{"s0": "c1"}]
    assert state.materialized_scene_responses == [scene_task.request_id]
    assert not state.accepted_product


@pytest.mark.parametrize("bad", ["required_units", "primary_group_id", "slide_id"])
def test_brief_contract_errors_do_not_consume_or_reissue(tmp_path, bad):
    state = case(tmp_path)
    task = state.task_queue[0]
    payload = brief()
    payload[bad] = [] if bad == "required_units" else "unknown"
    errors = runner.consume_response(state, response(task, payload))
    assert errors
    assert task.status == "pending" and state.tasks_used == 0
    assert len(state.task_queue) == 3


def test_unmapped_concepts_are_rejected_before_commit(tmp_path):
    state = case(tmp_path)
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    task = next(t for t in state.task_queue if t.kind == ModelTaskKind.VISUAL_CONCEPT_GENERATION)
    payload = concepts()
    payload["concepts"][0]["source_to_visual_mapping"] = []
    errors = runner.consume_response(state, response(task, payload))
    assert any("CONCEPT_FIDELITY_INVALID" in e for e in errors)
    assert task.status == "pending" and state.tasks_used == 1
    assert not state.slide_candidates


@pytest.mark.parametrize("bad", ["duplicate", "missing", "unknown_ref", "cosmetic_duplicate"])
def test_concept_set_integrity_is_preflighted(tmp_path, bad):
    state = case(tmp_path)
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    task = next(t for t in state.task_queue if t.kind == ModelTaskKind.VISUAL_CONCEPT_GENERATION)
    payload = concepts()
    if bad == "duplicate":
        payload["concepts"][1]["concept_id"] = "c1"
    elif bad == "missing":
        payload["concepts"].pop()
    elif bad == "unknown_ref":
        payload["concepts"][0]["source_to_visual_mapping"][0]["source_unit_ids"] = ["invented"]
    else:
        payload["concepts"][1]["group_topology"] = payload["concepts"][0]["group_topology"]
    assert runner.consume_response(state, response(task, payload))
    assert task.status == "pending" and state.tasks_used == 1


def test_dispatch_prefers_downstream_work_but_keeps_outstanding_immutable(tmp_path):
    root = tmp_path / "case/run"
    state = case(root)
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    order = dispatch.next_work(tmp_path)
    assert order["kind"] == "visual_concept_generation"
    work = dispatch.read(Path(order["work_order"]))
    assert "concepts" in work["payload_schema"]["required"]
    before = order["request_id"]
    assert dispatch.next_work(tmp_path)["request_id"] == before


def test_scene_work_order_schema_has_required_scene_plan_wrapper(tmp_path):
    root = tmp_path / "case/run"
    state = case(root, 1)
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    task = next(t for t in state.task_queue if t.kind == ModelTaskKind.VISUAL_CONCEPT_GENERATION)
    runner.advance_run(state, [response(task, concepts())])
    order = dispatch.next_work(tmp_path)
    assert order["kind"] == "scene_generation"
    work = dispatch.read(Path(order["work_order"]))
    assert "scene_plan" in work["payload_schema"]["required"]


@pytest.mark.parametrize("name", ["slide-1.png", "slide-01.png", "slide-001.png"])
def test_source_image_resolves_actual_numeric_page_name(tmp_path, name):
    case(tmp_path)
    path = tmp_path / "source_render" / name
    path.parent.mkdir()
    path.write_bytes(b"source-image")
    assert source_slide_images(tmp_path, "s0") == [str(path.resolve())]


def test_source_image_does_not_guess_between_different_images(tmp_path):
    case(tmp_path)
    root = tmp_path / "source_render"
    root.mkdir()
    (root / "slide-1.png").write_bytes(b"first")
    (root / "slide-01.png").write_bytes(b"other")
    with pytest.raises(ValueError, match="SOURCE_IMAGE_AMBIGUOUS"):
        source_slide_images(tmp_path, "s0")


def test_source_image_refuses_stale_render_probe(tmp_path):
    case(tmp_path)
    source = dispatch.read(tmp_path / "source_ir.json")
    source["source_sha256"] = "new-source"
    write(tmp_path / "source_ir.json", source)
    write(tmp_path / "source_render/source_render_probe.json", {"candidate_hash": "old-source"})
    with pytest.raises(ValueError, match="SOURCE_RENDER_STALE"):
        source_slide_images(tmp_path, "s0")


def test_model_receives_same_content_frame_and_units_as_engine(tmp_path):
    from slides_cli.gpn.frame_contract import model_frame_context

    case(tmp_path)
    write(
        tmp_path / "template_profile.json",
        {
            "content_box": {"x": 1, "y": 2, "w": 8, "h": 4},
            "unit": "in",
            "verified": True,
            "source": "test_only_verified_frame",
        },
    )
    frame = model_frame_context(tmp_path)
    assert frame["verified"]
    assert frame["content_box_pt"] == {"x": 72, "y": 144, "w": 576, "h": 288}
    assert "NOT percent of canvas" in frame["zone_coordinates"]
    assert (
        frame["profile_sha256"]
        == hashlib.sha256((tmp_path / "template_profile.json").read_bytes()).hexdigest()
    )


def test_comparison_includes_slides_without_concepts(tmp_path):
    state = case(tmp_path)
    state.slide_candidates = {"s0": ["c1", "c2"]}
    ok, missing = comparison_coverage(state, [])
    assert not ok and set(missing) == {"s0", "s1", "s2"}


def test_partial_frontier_waits_for_briefs_instead_of_blocking(tmp_path):
    state = case(tmp_path)
    state.stage = "candidate_selection"
    runner._advance_build_stages(state, [])
    assert state.stage == "build" and state.terminal_reason == ""
    assert not any(t.kind == ModelTaskKind.CANDIDATE_SELECTION for t in state.task_queue)


def test_audit_does_not_charge_313_historical_scenes_to_budget(tmp_path):
    root = tmp_path / "case/run"
    state = case(root, 1).model_dump(mode="json")
    pending = state["task_queue"][0]
    history = [dict(pending, request_id=f"history-{i}", status="rejected") for i in range(313)]
    state.update(
        task_queue=history + [pending],
        tasks_used=42,
        consumed_request_ids=[f"done-{i}" for i in range(42)],
    )
    state["budget"]["max_model_tasks"] = 408
    write(root / "run_state.json", state)
    audit = dispatch.audit_progress(tmp_path)["cases"][0]
    assert audit["historical_tasks"] == 314
    assert audit["unreserved_model_capacity"] == 365
    assert audit["counter_consistent"]


def test_review_bundle_keeps_successful_candidate_and_normalizes_absolute_packets(tmp_path):
    batch = tmp_path / "batch"
    root = batch / "case/run"
    state = case(root, 1).model_dump(mode="json")
    candidate = root / "scenes/slide-proof/candidate.pptx"
    candidate.parent.mkdir(parents=True)
    candidate.write_bytes(b"test-only bytes; bundle inclusion test, not a native gate")
    png = candidate.parent / "render/page-1.png"
    png.parent.mkdir()
    png.write_bytes(b"test-only image bytes")
    state["deck_candidates"] = [
        {
            "deck_key": "slide-proof",
            "eligible": True,
            "candidate.pptx": str(candidate.relative_to(root)),
            "sha256": hashlib.sha256(candidate.read_bytes()).hexdigest(),
        }
    ]
    write(root / "run_state.json", state)
    result = dispatch.bundle_review(batch, TOOL.parents[1])
    with zipfile.ZipFile(result["review_bundle"]) as archive:
        names = archive.namelist()
        assert "cases/case/scenes/slide-proof/candidate.pptx" in names
        assert "cases/case/scenes/slide-proof/render/page-1.png" in names
        assert "cases/case/model_requests/brief-s0.json" in names
        assert not any("/Users/" in n for n in names)


def test_case_fairness_precedes_phase_priority(tmp_path):
    root1, root2 = tmp_path / "old/run", tmp_path / "new/run"
    first = case(root1)
    case(root2)
    runner.advance_run(first, [response(first.task_queue[0], brief())])
    write(
        tmp_path / ".harness_dispatch.json",
        {"sequence": 1, "served": {str(root1): 1, str(root2): 0}},
    )
    assert dispatch.next_work(tmp_path)["run_dir"] == str(root2)


def own_library(project, text="Own example content"):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "DejaVu Sans"
            run.font.size = Pt(18)
    path = project / "slide_examples/own.pptx"
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
    return path


def test_owned_example_cache_changes_when_source_bytes_change(tmp_path):
    from slides_cli.gpn.owned_examples import select_owned_examples

    own_library(tmp_path)
    state = case(tmp_path / "run")
    slide = SemanticDeck.model_validate(
        dispatch.read(Path(state.run_dir) / "inputs/semantic_deck.json")
    ).slides[0]
    first = select_owned_examples(tmp_path, slide)
    assert first and first[0]["source_relative_path"] == "slide_examples/own.pptx"
    own_library(tmp_path, "Different own example revision")
    second = select_owned_examples(tmp_path, slide)
    assert second[0]["source_sha256"] != first[0]["source_sha256"]
    assert second[0]["title_snippet"] == "Different own example revision"


def test_foreign_example_is_refused_before_rendering(tmp_path, monkeypatch):
    from slides_cli.gpn.owned_examples import render_owned_examples
    from slides_cli.gpn.reference_guard import ForbiddenReferenceError

    own_library(tmp_path)
    foreign = tmp_path / "agent-slides/examples/example.pptx"
    foreign.parent.mkdir(parents=True)
    foreign.write_bytes(b"forbidden reference")

    def unexpected(*args, **kwargs):
        raise AssertionError("foreign example must never reach renderer")

    monkeypatch.setattr("slides_cli.gpn.render.render_source_deck", unexpected)
    with pytest.raises(ForbiddenReferenceError):
        render_owned_examples(
            tmp_path,
            [
                {
                    "source_path": str(foreign),
                    "source_sha256": "x",
                    "slide_id": "slide0:ppt/slides/slide1.xml",
                }
            ],
        )


def test_owned_example_png_is_real_and_cache_is_reused(tmp_path, monkeypatch):
    import shutil

    from slides_cli.gpn.owned_examples import render_owned_examples, select_owned_examples
    from slides_cli.gpn.render import render_source_deck

    if not shutil.which("soffice") or not shutil.which("pdftoppm"):
        pytest.skip("native input-render integration requires LibreOffice and pdftoppm")
    library = own_library(tmp_path)
    source_hash = hashlib.sha256(library.read_bytes()).hexdigest()
    state = case(tmp_path / "run")
    slide = SemanticDeck.model_validate(
        dispatch.read(Path(state.run_dir) / "inputs/semantic_deck.json")
    ).slides[0]
    selected = select_owned_examples(tmp_path, slide)
    calls = []

    def tracked(*args, **kwargs):
        calls.append(str(args[0]))
        return render_source_deck(*args, **kwargs)

    monkeypatch.setattr("slides_cli.gpn.render.render_source_deck", tracked)
    first = render_owned_examples(tmp_path, selected)
    assert not first["issues"], first
    image = Path(first["required_images"][0])
    assert image.read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
    second = render_owned_examples(tmp_path, selected)
    assert second == first
    assert len(calls) == 1
    assert hashlib.sha256(library.read_bytes()).hexdigest() == source_hash


def test_concept_request_receives_only_actual_owned_example_ids(tmp_path):
    root = tmp_path / "run"
    own_library(root)
    state = case(root)
    runner.advance_run(state, [response(state.task_queue[0], brief())])
    task = next(t for t in state.task_queue if t.kind == ModelTaskKind.VISUAL_CONCEPT_GENERATION)
    packet = dispatch.read(root / task.packet)
    assert packet["approved_example_refs"]
    assert all(ref.startswith("own.pptx:") for ref in packet["approved_example_refs"])
