"""Shared synthetic fixture builders for GPN tests (spec §21).

These fixtures validate engineering properties only; they never stand in for a
real corporate GPN master (test-only profile).
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches, Pt


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    return prs


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_textbox(slide, text: str, *, left=1.0, top=1.0, width=6.0, height=1.5, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    return box


def add_mixed_runs(slide):
    """Cyrillic, currency, percent, minus, subscript, hyperlink, soft break."""
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tf = box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    r1 = p0.add_run()
    r1.text = "Налоговая нагрузка: 12,5%"
    r2 = p0.add_run()
    r2.text = " ₽"
    p1 = tf.add_paragraph()
    r3 = p1.add_run()
    r3.text = "−42"
    r4 = p1.add_run()
    r4.text = "2"
    r4.font._rPr.set("baseline", "-25000")  # subscript
    p2 = tf.add_paragraph()
    r5 = p2.add_run()
    r5.text = "см. "
    r5.hyperlink.address = "https://example.invalid/doc"
    r6 = p2.add_run()
    r6.text = "пояснение"
    p3 = tf.add_paragraph()
    r7 = p3.add_run()
    r7.text = "строка один"
    from lxml import etree as _etree
    from pptx.oxml.ns import qn as _qn

    r7._r.addnext(_etree.Element(_qn("a:br")))
    r8 = p3.add_run()
    r8.text = "строка два"
    return box


def add_fields_and_footnotes(slide):
    """Fields (slidenum, date) with cached text, footnote marker, soft break."""
    from lxml import etree as _etree
    from pptx.oxml.ns import qn as _qn

    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tf = box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r1 = p0.add_run()
    r1.text = "Текст со сноской"

    # superscript footnote marker
    r_fn = p0.add_run()
    r_fn.text = "1"
    r_fn.font._rPr.set("baseline", "30000")

    # field: slide number
    fld_slidenum = _etree.SubElement(p0._p, _qn("a:fld"))
    fld_slidenum.set("id", "{11111111-1111-1111-1111-111111111111}")
    fld_slidenum.set("type", "slidenum")
    t_node = _etree.SubElement(fld_slidenum, _qn("a:t"))
    t_node.text = "1"

    # field: date-time
    fld_date = _etree.SubElement(p0._p, _qn("a:fld"))
    fld_date.set("id", "{22222222-2222-2222-2222-222222222222}")
    fld_date.set("type", "datetime1")
    t_node2 = _etree.SubElement(fld_date, _qn("a:t"))
    t_node2.text = "05.10.2026"

    # soft break
    p1 = tf.add_paragraph()
    r2 = p1.add_run()
    r2.text = "строка один"
    r2._r.addnext(_etree.Element(_qn("a:br")))
    r3 = p1.add_run()
    r3.text = "строка два"

    return box


def add_table(slide):
    """3x4 table with header row, horizontal + vertical merge, blank/0/NA cells."""
    table_shape = slide.shapes.add_table(3, 4, Inches(1), Inches(3), Inches(8), Inches(2))
    table = table_shape.table
    headers = ["Регион", "Выручка", "Доля", "Δ"]
    for c, text in enumerate(headers):
        table.cell(0, c).text = text
    table.cell(1, 0).text = "Север"
    table.cell(1, 1).text = "120,5"
    table.cell(1, 2).text = "0"
    table.cell(1, 3).text = "н/д"
    table.cell(2, 0).text = "Юг"
    table.cell(2, 1).text = ""
    table.cell(2, 2).text = "33,3%"
    table.cell(2, 3).text = "1"
    # horizontal merge of row 1, cols 1..2 via gridSpan
    _set_grid_span(table.cell(1, 1), 2)
    _mark_hmerge(table.cell(1, 2))
    # vertical merge of col 0 rows 1..2
    _set_row_span(table.cell(1, 0), 2)
    _mark_vmerge(table.cell(2, 0))
    return table_shape


def _tc_pr(cell):
    from pptx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.find(qn("a:tcPr"))
    if tc_pr is None:
        from lxml import etree

        tc_pr = etree.SubElement(tc, qn("a:tcPr"))
    return tc_pr


def _set_grid_span(cell, span: int) -> None:
    cell._tc.set("gridSpan", str(span))


def _set_row_span(cell, span: int) -> None:
    cell._tc.set("rowSpan", str(span))


def _mark_hmerge(cell) -> None:
    cell._tc.set("hMerge", "1")


def _mark_vmerge(cell) -> None:
    cell._tc.set("vMerge", "1")


def add_column_chart(slide):
    data = CategoryChartData()
    data.categories = ["2023", "2024", "2023", "2025"]
    data.add_series("План", (10.5, 0, 12.25, None))
    data.add_series("Факт", (9.75, 11.0, 13.5, 8.0))
    return slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(3.5),
        Inches(7),
        Inches(3.5),
        data,
    )


def add_xy_chart(slide):
    data = XyChartData()
    s1 = data.add_series("T1")
    s1.add_data_point(1.0, 2.5)
    s1.add_data_point(2.0, 3.5)
    s2 = data.add_series("T2")
    s2.add_data_point(1.0, 4.5)
    s2.add_data_point(3.0, 1.5)
    return slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        Inches(1),
        Inches(3.5),
        Inches(7),
        Inches(3.5),
        data,
    )


def add_group_with_nonuniform_scale(slide):
    """Group with non-uniform scale, rotation, flip and a nested child."""
    group = slide.shapes.add_group_shape()
    child1 = group.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(2), Inches(1)
    )
    child1.fill.solid()
    child1.fill.fore_color.rgb = RGBColor(0x00, 0x45, 0x96)
    child1.text_frame.text = "A"
    nested = group.shapes.add_group_shape()
    child2 = nested.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(0), Inches(1.2), Inches(1.5), Inches(0.8)
    )
    child2.text_frame.text = "B"
    _force_group_xfrm(group, scale_x=1.5, scale_y=0.75, rotation=15 * 60000, flip_h=True)
    _force_group_xfrm(nested, scale_x=0.8, scale_y=1.2, rotation=-10 * 60000, flip_v=True)
    return group


def _force_group_xfrm(
    group, *, scale_x: float, scale_y: float, rotation: int, flip_h=False, flip_v=False
):
    from pptx.oxml.ns import qn

    grp_sp_pr = group._element.find(qn("p:grpSpPr"))
    xfrm = grp_sp_pr.find(qn("a:xfrm"))
    off = xfrm.find(qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    ch_off = xfrm.find(qn("a:chOff"))
    ch_ext = xfrm.find(qn("a:chExt"))
    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
    ext.set("cx", str(int(cx * scale_x)))
    ext.set("cy", str(int(cy * scale_y)))
    ch_ext.set("cx", str(cx))
    ch_ext.set("cy", str(cy))
    ch_off.set("x", off.get("x"))
    ch_off.set("y", off.get("y"))
    xfrm.set("rot", str(rotation))
    if flip_h:
        xfrm.set("flipH", "1")
    if flip_v:
        xfrm.set("flipV", "1")


def add_zero_extent_connector(slide):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(4), Inches(6), Inches(4)
    )
    line.line.color.rgb = RGBColor(0x00, 0x00, 0x00)
    # force zero height bounding box (horizontal line)
    line._element.spPr.xfrm.ext.set("cy", "0")
    return line


def add_internal_link_and_notes(prs):
    s1 = blank_slide(prs)
    add_textbox(s1, "Первый слайд")
    s1.notes_slide.notes_text_frame.text = "Заметки к первому слайду"
    s2 = blank_slide(prs)
    box = add_textbox(s2, "Второй слайд")
    _add_internal_hyperlink(box, s1)
    s2._element.set("show", "0")
    return prs, 2


def _add_internal_hyperlink(shape, target_slide):
    from lxml import etree
    from pptx.oxml.ns import qn

    r = shape.text_frame.paragraphs[0].runs[0]._r
    rPr = r.find(qn("a:rPr"))
    if rPr is None:
        rPr = etree.SubElement(r, qn("a:rPr"))
        r.insert(0, rPr)
    hlink = etree.SubElement(rPr, qn("a:hlinkClick"))
    rid = shape.part.relate_to(
        target_slide.part,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
        is_external=False,
    )
    hlink.set(qn("r:id"), rid)


def build_import_fixture(path: Path) -> Path:
    """Full fixture deck: mixed runs, equal text, table, chart, group, connector."""
    prs = new_deck()
    s1 = blank_slide(prs)
    add_mixed_runs(s1)
    add_textbox(s1, "Одинаковый текст", left=1, top=3, width=4, height=0.6, size=14)
    add_textbox(s1, "Одинаковый текст", left=1, top=4, width=4, height=0.6, size=14)
    add_table(s1)
    s2 = blank_slide(prs)
    add_column_chart(s2)
    add_zero_extent_connector(s2)
    add_group_with_nonuniform_scale(s2)
    add_unknown_graphic_data(s2)
    s3 = blank_slide(prs)
    add_xy_chart(s3)
    prs.save(str(path))
    return path


def add_unknown_graphic_data(slide):
    """GraphicFrame with an unsupported graphicData URI (must be detected)."""
    from lxml import etree
    P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
    A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
    R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    xml = (
        f'<p:graphicFrame xmlns:p="{P_NS}" xmlns:a="{A_NS}" '
        f'xmlns:r="{R_NS}">'
        "<p:nvGraphicFramePr>"
        '<p:cNvPr id="9999" name="UnknownDiagram"/>'
        "<p:cNvGraphicFramePr/>"
        "<p:nvPr/>"
        "</p:nvGraphicFramePr>"
        "<p:xfrm><a:off x=\"1000000\" y=\"1000000\"/>"
        "<a:ext cx=\"2000000\" cy=\"1000000\"/></p:xfrm>"
        "<a:graphic><a:graphicData "
        'uri="http://schemas.microsoft.com/office/drawing/2017/diagram">'
        "<dgm:relIds xmlns:dgm=\"http://schemas.openxmlformats.org/drawingml/2006/diagram\" "
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:dm="rId1" r:lo="rId2" r:qs="rId3"/>'
        "</a:graphicData></a:graphic>"
        "</p:graphicFrame>"
    )
    frame_el = etree.fromstring(xml)
    slide.shapes._spTree.append(frame_el)
    return frame_el
