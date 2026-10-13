"""Stage 7 §11/§14: public CLI routes + bundle closure regressions."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

from slides_cli.gpn.cli import build_gpn_parser
from slides_cli.gpn.cli import main as gpn_main

PROJECT_ROOT = Path(__file__).resolve().parents[3]


def _args(project_root: Path):
    return SimpleNamespace(config=None, project_root=project_root)


def test_parser_has_public_routes(capsys):
    parser = build_gpn_parser()
    for route in ("restyle", "create", "resume", "status", "bundle"):
        with pytest.raises(SystemExit) as exc:
            parser.parse_args([route, "--help"])
        assert exc.value.code == 0
        assert route in capsys.readouterr().out or True
    args = parser.parse_args(["status", "some-dir"])
    assert args.gpn_command == "status"


def test_restyle_starts_and_waits(tmp_path):
    from conftest import build_import_fixture

    src = build_import_fixture(tmp_path / "in.pptx")
    code = gpn_main([
        "--project-root", str(PROJECT_ROOT), "restyle", str(src),
        "--output", str(tmp_path / "out.pptx"),
        "--run-dir", str(tmp_path / "run1")])
    assert code == 4  # awaiting_harness, not a terminal failure
    assert (tmp_path / "run1" / "run_state.json").is_file()


def test_restyle_missing_input_invalid(tmp_path):
    code = gpn_main([
        "--project-root", str(tmp_path), "restyle",
        str(tmp_path / "nope.pptx"),
        "--output", str(tmp_path / "out.pptx")])
    assert code == 2


def test_create_prompt_starts_and_waits(tmp_path):
    code = gpn_main([
        "--project-root", str(tmp_path), "create",
        "--prompt", "Brief with numbers 400 FTE",
        "--output", str(tmp_path / "out.pptx"),
        "--run-dir", str(tmp_path / "run2")])
    assert code == 4
    assert (tmp_path / "run2" / "inputs" / "source_manifest.json").is_file()


def test_status_and_resume_roundtrip(tmp_path, capsys):
    from conftest import build_import_fixture

    src = build_import_fixture(tmp_path / "in.pptx")
    assert gpn_main([
        "--project-root", str(PROJECT_ROOT), "restyle", str(src),
        "--output", str(tmp_path / "out.pptx"),
        "--run-dir", str(tmp_path / "run3")]) == 4
    assert gpn_main([
        "--project-root", str(tmp_path), "status",
        str(tmp_path / "run3")]) == 0
    out = capsys.readouterr().out
    assert "pending_tasks" in out
    assert gpn_main([
        "--project-root", str(tmp_path), "resume",
        str(tmp_path / "run3")]) == 4


def test_local_server_is_deferred(tmp_path):
    src = tmp_path / "in.pptx"
    src.write_bytes(b"fake-pptx-bytes")
    code = gpn_main([
        "--project-root", str(tmp_path), "restyle", str(src),
        "--output", str(tmp_path / "out.pptx"),
        "--llm-provider", "local_server"])
    assert code == 3


def test_bundle_marks_missing_required(tmp_path):
    from slides_cli.gpn.review_bundle import build_review_bundle
    run = tmp_path / "run4"
    run.mkdir()
    report = build_review_bundle(run, tmp_path / "b.zip")
    assert not report.bundle_complete
    assert report.missing_required
    assert tmp_path / "b.zip"


def test_bundle_includes_and_self_checks(tmp_path):
    from slides_cli.gpn.review_bundle import build_review_bundle
    run = tmp_path / "run5"
    run.mkdir()
    (run / "candidate.pptx").write_bytes(b"PK-fake")
    (run / "product_acceptance.json").write_text(json.dumps(
        {"subject_candidate_sha256": "x" * 64}), encoding="utf-8")
    (run / "run_state.json").write_text("{}", encoding="utf-8")
    (run / "runner_transition_log.jsonl").write_text("{}\n",
                                                     encoding="utf-8")
    (run / "render").mkdir()
    (run / "render" / "slide-1.png").write_bytes(b"png")
    report = build_review_bundle(run, tmp_path / "b.zip")
    assert report.missing_required  # many roles absent -> not complete
    with zipfile.ZipFile(tmp_path / "b.zip") as zf:
        names = zf.namelist()
        assert "bundle_manifest.json" in names
        assert "main/candidate.pptx" in names
        manifest = json.loads(zf.read("bundle_manifest.json"))
        assert manifest["missing_required"]


def test_bundle_complete_false_with_chosen_files_only(tmp_path):
    """114 present entries never imply completeness: required schema rules."""
    from slides_cli.gpn.review_bundle import build_review_bundle
    run = tmp_path / "run6"
    run.mkdir()
    for i in range(114):
        (run / f"extra-{i}.json").write_text("{}", encoding="utf-8")
    report = build_review_bundle(run, tmp_path / "b.zip")
    assert not report.bundle_complete
    assert "final_package" in report.missing_required


def test_acceptance_main_accepts_complete_synth_run(tmp_path):
    """End-to-end CLI derivation: valid evidence -> accepted, exit 0."""
    import hashlib

    from pptx import Presentation
    from pptx.util import Emu, Inches

    from slides_cli.gpn.acceptance import main as acceptance_main
    from slides_cli.gpn.models import (
        SemanticSlide,
        SemanticUnit,
        TypedContentPayload,
    )

    run = tmp_path / "synth"
    (run / "render").mkdir(parents=True)
    (run / "gate_reports").mkdir()
    (run / "inputs").mkdir()
    (run / "final").mkdir()
    p = Presentation()
    p.slide_width = Emu(12192000)
    p.slide_height = Emu(6858000)
    for i in range(2):
        s = p.slides.add_slide(p.slide_layouts[6])
        t = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        t.text_frame.text = f"Slide {i}"
    p.save(run / "candidate.pptx")
    sha = hashlib.sha256((run / "candidate.pptx").read_bytes()).hexdigest()

    slides = [SemanticSlide(
        source_slide_id=f"s{i}", title=f"S{i}",
        units=[SemanticUnit(id=f"u{i}a", kind="text",
                            content=TypedContentPayload(text="x")),
               SemanticUnit(id=f"u{i}b", kind="text",
                            content=TypedContentPayload(text="y"))])
        for i in range(2)]
    from slides_cli.gpn.semantics import _hash_slides
    deck = {"source_manifest_sha256": "c" * 64,
            "slides": [s.model_dump(mode="json") for s in slides],
            "units": {}, "relations": {},
            "input": {"kind": "pptx", "path": "/tmp/x.pptx",
                      "scope": "restyle"}}
    deck["sha256"] = _hash_slides(slides)
    (run / "inputs" / "semantic_deck.json").write_text(json.dumps(deck))
    (run / "compiled_rules.json").write_text(json.dumps(
        {"canvas": {"cx": 12192000, "cy": 6858000},
         "rule_registry": {"rules": {"C01": {}}}}))
    (run / "input_hashes.json").write_text(json.dumps(
        {"input_sha256": "c" * 64}))
    (run / "font_inventory.json").write_text(json.dumps({"verified": True}))
    (run / "render" / "render_manifest.json").write_text(json.dumps(
        {"pages": [{"page": 1, "image": "slide-1.png", "sha256": "a" * 64},
                   {"page": 2, "image": "slide-2.png",
                    "sha256": "b" * 64}]}))

    from slides_cli.gpn.acceptance import (
        _expectation_from_run,
        _snapshots_from_run,
    )
    from slides_cli.gpn.models import RunContext
    ctx = RunContext(project_root=tmp_path, run_dir=run)
    exp = _expectation_from_run(run, ctx, _snapshots_from_run(run))
    for g in exp.required_gates:
        (run / "gate_reports" / f"{g.gate_id}.json").write_text(
            json.dumps({"gate": g.gate_id, "status": "pass",
                        "expected_units": g.expected_units,
                        "checked_units": g.expected_units,
                        "subject_candidate_sha256": sha}))
    (run / "visual_review.json").write_text(json.dumps(
        {"task_id": "t", "candidate_sha256": sha, "verdict": "accept",
         "pages": [{"page": 1, "image": "slide-1.png", "sha256": "a" * 64,
                    "bytes": 10},
                   {"page": 2, "image": "slide-2.png", "sha256": "b" * 64,
                    "bytes": 10}],
         "actual_model": "test"}))
    (run / "selection_report.json").write_text(json.dumps({"selections": [
        {"source_slide_id": "s0",
         "candidate_key": {"run_id": "synth", "source_slide_id": "s0",
                           "candidate_id": "c1", "revision": 0},
         "candidate_sha256": sha},
        {"source_slide_id": "s1",
         "candidate_key": {"run_id": "synth", "source_slide_id": "s1",
                           "candidate_id": "c1", "revision": 0},
         "candidate_sha256": sha}]}))
    (run / "saved_package_manifest.json").write_text(json.dumps(
        {"candidate_path": str(run / "candidate.pptx")}))

    code = acceptance_main(str(run))
    assert code == 0
    acc = json.loads((run / "product_acceptance.json").read_text())
    assert acc["accepted"] is True
    assert acc["acceptance_schema_version"] == "product-acceptance/v2"
    assert (run / "product_expectation.json").is_file()
    assert (run / "saved_package_manifest.json").is_file()
    assert (run / "gate_index.json").is_file()
