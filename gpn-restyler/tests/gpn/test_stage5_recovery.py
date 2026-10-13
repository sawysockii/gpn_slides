"""Stage 5 recovery regression tests (failure modes of pilot-ntc 82798350).

Every test asserts a real invariant on serialized bytes — never a mirror of
the implementation. Fixtures are small synthetic decks built in tmp_path
with the same failure mode as the bad candidate (roundRect + off-palette
fill + inherited effectRef + synthetic bold + colorless runs); the user
confidential binaries are never committed. Normative values flow from the
loaded ontology snapshot, never from literals repeated here.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

PROJECT_ROOT = Path(__file__).resolve().parents[3]

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _rules():
    from slides_cli.gpn.compiler import compile_ontology
    from slides_cli.gpn.ontology import resolve_ontology_sources

    sources = resolve_ontology_sources(PROJECT_ROOT)
    assert sources.primary_json is not None
    return compile_ontology(Path(sources.primary_json), None, sources=sources)


def _bad_deck(path: Path) -> Path:
    """Synthetic deck with the bad-candidate failure mode."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Decorative rounded card with off-palette fill (C01/C03).
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1),
        Inches(4), Inches(2))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0x0C, 0x53, 0x9F)
    # Text with synthetic bold and no explicit color (C21/C22/foreground).
    box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Recovery probe"
    run.font.bold = True
    prs.save(str(path))
    return path


def _fixed_deck(path: Path) -> Path:
    """Same content through the normative path: rect, no fill, no bold."""
    from slides_cli.gpn import opc_adapter as adapter

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    card.fill.background()
    adapter.strip_shape_effect_inheritance(card)
    box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Recovery probe"
    run.font.bold = False
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    adapter.strip_shape_effect_inheritance(box)
    prs.save(str(path))
    return path


def _check(path: Path, rules):
    from slides_cli.gpn.rules import check_saved_authoring_rules

    return check_saved_authoring_rules(path, rules)


# ---------------------------------------------------------------------------
# §9: actual bad-candidate failure mode fails on reopened bytes
# ---------------------------------------------------------------------------

def test_saved_authoring_fails_on_bad_fixture(tmp_path: Path):
    rules = _rules()
    deck = _bad_deck(tmp_path / "bad.pptx")
    report = _check(deck, rules)
    assert report["slides_checked"] == 1
    assert report["hard_fail"] is True
    assert report["rules"]["C01"]["result"] == "fail"
    assert report["rules"]["C03"]["result"] == "fail"
    assert report["rules"]["C21_C22"]["result"] == "fail"
    assert report["rules"]["surface_foreground"]["result"] == "fail"


def test_saved_authoring_passes_on_fixed_fixture(tmp_path: Path):
    rules = _rules()
    deck = _fixed_deck(tmp_path / "fixed.pptx")
    report = _check(deck, rules)
    assert report["hard_fail"] is False
    for rid, info in report["rules"].items():
        assert info["result"] == "pass", f"{rid}: {info['violations'][:2]}"


# ---------------------------------------------------------------------------
# §9: inherited effectRef fails even with no direct effectLst
# ---------------------------------------------------------------------------

def test_c03_effect_ref_without_effect_list_fails():
    from slides_cli.gpn.rules import RuleEvaluationContext, evaluate_rule

    rules = _rules()
    raw = ("<p:sp><p:spPr><a:prstGeom prst=\"rect\"/>"
           "<p:style><a:effectRef idx=\"2\"/></p:style>"
           "</p:spPr></p:sp>")
    ctx = RuleEvaluationContext(
        subject_kind="target", compiled_rules=rules, raw_package_xml=raw)
    coverage = evaluate_rule("C03", ctx)
    assert coverage.result == "fail"
    assert "effectRef" in coverage.reason


def test_c03_clean_shape_passes():
    from slides_cli.gpn.rules import RuleEvaluationContext, evaluate_rule

    rules = _rules()
    raw = ("<p:sp><p:spPr><a:prstGeom prst=\"rect\"/>"
           "<a:effectLst/></p:spPr></p:sp>")
    ctx = RuleEvaluationContext(
        subject_kind="target", compiled_rules=rules, raw_package_xml=raw)
    coverage = evaluate_rule("C03", ctx)
    assert coverage.result == "pass"


# ---------------------------------------------------------------------------
# §9: fresh library shapes carry no shadow after normalization
# ---------------------------------------------------------------------------

def test_fresh_shape_effect_inheritance_stripped(tmp_path: Path):
    from slides_cli.gpn import opc_adapter as adapter

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    adapter.strip_shape_effect_inheritance(shape)
    path = tmp_path / "fresh.pptx"
    prs.save(str(path))
    with zipfile.ZipFile(path) as archive:
        names = [n for n in archive.namelist()
                 if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        assert names, "no slide part saved"
        xml = archive.read(names[0]).decode("utf-8")
    assert "effectRef" not in xml


# ---------------------------------------------------------------------------
# §9: late writer override is detected on final bytes
# ---------------------------------------------------------------------------

def test_late_writer_bold_override_detected(tmp_path: Path):
    rules = _rules()
    deck = _fixed_deck(tmp_path / "override.pptx")
    # Rewrite the package with the mutated slide part (no duplicate entry:
    # a late writer override must be visible to the next reader).
    mutated = tmp_path / "overridden.pptx"
    with zipfile.ZipFile(deck, "r") as src:
        names = [n for n in src.namelist()
                 if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        assert names, "no slide part saved"
        target = names[0]
        raw = src.read(target)
        root = etree.fromstring(raw)
        ns = {"a": A_NS}
        rprs = root.findall(".//a:rPr", ns)
        assert rprs, "no runs to override"
        rprs[0].set("b", "1")
        patched = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                 standalone=True)
        with zipfile.ZipFile(mutated, "w",
                             compression=zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = patched if item.filename == target else src.read(
                    item.filename)
                dst.writestr(item, data)
    report = _check(mutated, rules)
    assert report["hard_fail"] is True
    assert report["rules"]["C21_C22"]["result"] == "fail"


# ---------------------------------------------------------------------------
# §9: proposal content never selects Contacts furniture
# ---------------------------------------------------------------------------

def test_layout_selection_rejects_contacts():
    from slides_cli.gpn.opc_adapter import blank_slide_layout

    class _Stub:
        def __init__(self, name, n_ph):
            self.name = name
            self._n = n_ph

        @property
        def placeholders(self):
            return [object()] * self._n

    class _Prs:
        slide_layouts = [
            _Stub("Контакты", 3),
            _Stub("Title and Content", 2),
            _Stub("Blank", 0),
        ]

    chosen = blank_slide_layout(_Prs())
    assert "контакт" not in (chosen.name or "").lower()
    assert "contact" not in (chosen.name or "").lower()


def test_layout_selection_prefers_blank_over_contacts():
    from slides_cli.gpn.opc_adapter import blank_slide_layout

    class _Stub:
        def __init__(self, name, n_ph):
            self.name = name
            self._n = n_ph

        @property
        def placeholders(self):
            return [object()] * self._n

    class _Prs:
        # Contacts first (index 0): the old hardcoded index=6 bug class.
        slide_layouts = [_Stub("Контакты", 1), _Stub("Blank", 0)]

    chosen = blank_slide_layout(_Prs(), layout_index=0)
    assert (chosen.name or "") == "Blank"
