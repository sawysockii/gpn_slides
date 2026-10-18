"""Stage 5 composition + resolution tests (spec §§6, 7, 11.2, 13.1).

Synthetic slides, real compiled ontology snapshot, real Chromium measurement
through the shared bridge. No LLM calls; intents are hand-built valid
LayoutIntents (the harness path is covered by the pilot runs).
"""

from __future__ import annotations

import json
from functools import lru_cache
from pathlib import Path

import conftest as fx
import pytest

from slides_cli.gpn.composition import OPERATOR_BINDINGS, compile_composition
from slides_cli.gpn.layout import resolve_layout
from slides_cli.gpn.planner import LayoutIntent

PROJECT_ROOT = Path(__file__).resolve().parents[3]

TEST_PROFILE = {"content_box": {"x": 0.5, "y": 1.2, "w": 9.0, "h": 4.3},
                "unit": "in", "verified": True, "source": "test_template"}


@lru_cache(maxsize=1)
def _rules():
    from slides_cli.gpn.compiler import compile_ontology
    from slides_cli.gpn.ontology import resolve_ontology_sources

    sources = resolve_ontology_sources(PROJECT_ROOT)
    assert sources.primary_json is not None
    return compile_ontology(PROJECT_ROOT / sources.primary_json, None,
                            sources=sources)


@lru_cache(maxsize=1)
def _bundle():
    from slides_cli.gpn.typography import (
        build_font_inventory_from_discovery,
        discover_local_fonts,
        prepare_measurement_fonts,
    )

    report = discover_local_fonts(None)
    faces = {f for role in _rules().roles.values() for f in role.allowed_faces}
    inv = build_font_inventory_from_discovery(report, faces)
    assert inv.missing_faces == []
    return prepare_measurement_fonts(inv, "", profile_kind="production")


def _import_deck(slide_texts: list[str], tmp_path: Path, name: str = "src"):
    from slides_cli.gpn.assets import AssetStore
    from slides_cli.gpn.importer import import_deck
    from slides_cli.gpn.provenance import build_ledger

    prs = fx.new_deck()
    slide = fx.blank_slide(prs)
    for i, text in enumerate(slide_texts):
        fx.add_textbox(slide, text, left=1.0, top=1.0 + i * 1.2,
                       width=6.0, height=1.0, size=18)
    src = tmp_path / f"{name}.pptx"
    prs.save(str(src))
    store = AssetStore(tmp_path / "assets")
    deck, _ = import_deck(src, store)
    ledger = build_ledger(deck, decoration=[])
    return deck, ledger


def _intent(slide_id: str, refs: list[str], **kw) -> LayoutIntent:
    zones = kw.pop("zones", [{"zone_id": "z-all", "x_pct": 0.0, "y_pct": 0.0,
                              "w_pct": 100.0, "h_pct": 100.0, "role": "body",
                              "content_refs": list(refs)}])
    data = {"slide_id": slide_id, "candidate_id": kw.pop("candidate_id", "c0"),
            "content_refs": list(refs), "communication_job": "test",
            "reading_order": list(refs), "dominant_ref": refs[0] if refs else "",
            "zones": zones, "relations": kw.pop("relations", []),
            "operators": kw.pop("operators", []), "reference_ids": [],
            "novel_recipe": "", "uncertainties": [], "rationale": "test"}
    data.update(kw)
    return LayoutIntent.model_validate(data)


def _measurer(options: dict | None = None):
    from slides_cli.gpn.measurement import measure_components

    seen: list[dict] = []

    def run(request):
        seen.append(request.model_dump(mode="json"))
        out = measure_components(
            request, timeout_seconds=120.0, project_root=PROJECT_ROOT,
            use_cache=True,
            debug_output_dir=str(options.get("debug_dir", "")) if options else "")
        return out["response"]

    run.seen = seen
    return run


def _options(tmp_path: Path, **kw) -> dict:
    opts = {"project_root": PROJECT_ROOT,
            "font_allowlist_roots": ["/Users/wysockii/Library/Fonts/",
                                     "/System/Library/Fonts/"],
            "use_cache": True, "profile_kind": "production",
            "max_local_layout_attempts": 3,
            "debug_output_dir": str(tmp_path / "layout_debug")}
    opts.update(kw)
    return opts


# ---------------------------------------------------------------------------
# §6 operators
# ---------------------------------------------------------------------------

def test_operator_constraints_have_effect(tmp_path: Path) -> None:
    """Every supported operator adds at least one geometry constraint or
    native construct — never only a rationale line."""
    deck, ledger = _import_deck(["Альфа", "Бета"], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:2]
    base = _intent(slide.id, refs, candidate_id="base")
    base_tree = compile_composition(base, slide, ledger, _rules(), TEST_PROFILE)
    n_base = len(base_tree.constraints)
    zones = [
        {"zone_id": "z-a", "x_pct": 0.0, "y_pct": 0.0, "w_pct": 48.0,
         "h_pct": 100.0, "role": "body", "content_refs": refs[:1]},
        {"zone_id": "z-b", "x_pct": 52.0, "y_pct": 0.0, "w_pct": 48.0,
         "h_pct": 100.0, "role": "body", "content_refs": refs[1:]},
    ]
    relations = [{"from_ref": refs[0], "to_ref": refs[1],
                  "relation_kind": "connects", "evidence": []}]
    cases = {
        "align": {}, "juxtapose": {"zones": zones}, "contrast": {"zones": zones},
        "sequence": {}, "nest": {"zones": zones}, "frame": {"zones": zones},
        "connect": {"relations": relations}, "branch": {"relations": relations},
        "annotate": {}, "highlight": {}, "overlay": {"zones": zones},
        "repeat": {}, "facet": {},
    }
    assert set(cases) <= set(OPERATOR_BINDINGS)
    for op, extra in cases.items():
        op_refs = list(refs)
        if op in ("repeat", "facet"):
            # authorized_repeat needs a real second occurrence
            op_refs = refs + refs[:1]
            extra = {"zones": [
                {"zone_id": "z-a", "x_pct": 0.0, "y_pct": 0.0, "w_pct": 100.0,
                 "h_pct": 100.0, "role": "body", "content_refs": op_refs}]}
        intent = _intent(slide.id, op_refs, candidate_id=f"op-{op}",
                         operators=[op], **extra)
        tree = compile_composition(intent, slide, ledger, _rules(), TEST_PROFILE)
        assert tree.operator_coverage[op] == "implemented", (
            op, tree.operator_coverage, tree.issues)
        assert len(tree.constraints) > n_base, op


def test_new_operator_needs_binding(tmp_path: Path) -> None:
    """A normative operator without an implementation binding — and a made-up
    one — get needs_binding, never a fake pass."""
    deck, ledger = _import_deck(["Альфа", "Бета"], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:2]
    for op in ("aggregate", "teleport"):
        intent = _intent(slide.id, refs, candidate_id=f"nb-{op}", operators=[op])
        tree = compile_composition(intent, slide, ledger, _rules(), TEST_PROFILE)
        assert tree.operator_coverage[op] == "needs_binding"
        assert tree.issues


def test_forbidden_reference_plan_cache_skill(tmp_path: Path) -> None:
    """A plan with foreign reference lineage is refused at compile time —
    through index, packet, cache or skill paths alike (no fallback)."""
    deck, ledger = _import_deck(["Альфа"], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:1]
    for bad in ("gpn-restyler/examples/example1.pptx:slide0",
                "agent-slides/demo/deck.pptx:slide0",
                "https://example.invalid/sample.pptx"):
        intent = _intent(slide.id, refs, reference_ids=[bad])
        with pytest.raises(ValueError, match="REFERENCE/FORBIDDEN_SOURCE"):
            compile_composition(intent, slide, ledger, _rules(), TEST_PROFILE)
    ok_intent = _intent(slide.id, refs,
                        reference_ids=["slide_examples.pptx:slide5"])
    tree = compile_composition(ok_intent, slide, ledger, _rules(), TEST_PROFILE)
    assert tree.components


# ---------------------------------------------------------------------------
# §6.2 diversity without a fixed grid
# ---------------------------------------------------------------------------

def test_three_distinct_compositions(tmp_path: Path) -> None:
    """The same content compiles into three genuinely different topologies
    (side-by-side / stacked / overlay) with all refs preserved."""
    deck, ledger = _import_deck(["Первый", "Второй", "Третий"], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:3]
    h_zones = [
        {"zone_id": f"z{i}", "x_pct": float(i * 34), "y_pct": 0.0,
         "w_pct": 32.0, "h_pct": 100.0, "role": "body",
         "content_refs": [r]} for i, r in enumerate(refs)]
    v_zones = [
        {"zone_id": f"z{i}", "x_pct": 0.0, "y_pct": float(i * 34),
         "w_pct": 100.0, "h_pct": 32.0, "role": "body",
         "content_refs": [r]} for i, r in enumerate(refs)]
    o_zones = [
        {"zone_id": f"z{i}", "x_pct": float(i * 10), "y_pct": float(i * 10),
         "w_pct": 80.0, "h_pct": 80.0, "role": "body",
         "content_refs": [r]} for i, r in enumerate(refs)]
    trees = [compile_composition(
        _intent(slide.id, refs, candidate_id=f"div{i}", zones=zones),
        slide, ledger, _rules(), TEST_PROFILE)
        for i, zones in enumerate((h_zones, v_zones, o_zones))]
    orgs = [t.root.organization for t in trees]
    assert orgs == ["h-stack", "v-stack", "overlay"]
    for tree in trees:
        bound = sorted(b.source_object_id for b in tree.content_bindings)
        assert bound == sorted(refs)


def test_protected_regions_not_reallocated(tmp_path: Path) -> None:
    """Allocation never moves into the reserved title band: every component
    box stays inside the verified content box (y >= 1.2in here)."""
    deck, ledger = _import_deck(["Альфа", "Бета", "Гамма"], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:3]
    intent = _intent(slide.id, refs, candidate_id="prot")
    tree = compile_composition(intent, slide, ledger, _rules(), TEST_PROFILE)
    res = resolve_layout(tree, slide, ledger, _rules(), TEST_PROFILE,
                         _bundle(), _measurer(), _options(tmp_path))
    assert res.status == "resolved", res.unresolved
    top_emu = min(box["y"] for box in res.resolved_slide["geometry_emu"].values())
    assert top_emu >= int(round(1.2 * 914400)) - 9144


def test_no_hidden_font_shrink(tmp_path: Path) -> None:
    """Overload never triggers fit_text/autofit/size clamps: the measured
    request keeps exact role sizes and attempts use only allowed actions."""
    from slides_cli.gpn.layout import ALLOWED_ATTEMPTS
    from slides_cli.gpn.measured_placement import SOLVER_EVIDENCE_ACTIONS

    deck, ledger = _import_deck(["Перегруженный " * 400], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:1]
    tiny = {"content_box": {"x": 0.5, "y": 1.2, "w": 9.0, "h": 0.4},
            "unit": "in", "verified": True, "source": "test_template"}
    intent = _intent(slide.id, refs, candidate_id="overload")
    tree = compile_composition(intent, slide, ledger, _rules(), tiny)
    measurer = _measurer()
    res = resolve_layout(tree, slide, ledger, _rules(), tiny,
                         _bundle(), measurer, _options(tmp_path))
    assert res.status in ("infeasible", "needs_recompose",
                          "search_exhausted", "constraint_conflict")
    role_size = _rules().roles["body"].default_size_pt
    for seen in measurer.seen:
        for comp in seen["components"]:
            for para in comp["paragraphs"]:
                for run in para["runs"]:
                    assert run["style"]["size_pt"] == role_size
    for attempt in res.attempts:
        for change in attempt.changes:
            assert change["action"] in (
                *ALLOWED_ATTEMPTS, "initial_allocation", *SOLVER_EVIDENCE_ACTIONS
            )
        assert "overflow_y" in " ".join(attempt.hard_issues)
    # The solver evidence records measured variants only. It never carries a
    # run, a box or a type override that could shrink type, drop text or turn
    # on autofit: only measured width/height numbers and solver bookkeeping.
    forbidden = {"runs", "text", "size_pt", "box_emu", "box_pt", "autofit",
                 "font_scale", "scale", "dropped", "paragraphs"}

    def _assert_evidence_only(value):
        if isinstance(value, dict):
            assert not (forbidden & set(value)), value
            for item in value.values():
                _assert_evidence_only(item)
        elif isinstance(value, list):
            for item in value:
                _assert_evidence_only(item)

    for attempt in res.attempts:
        for change in attempt.changes:
            if change["action"] in SOLVER_EVIDENCE_ACTIONS:
                _assert_evidence_only(change)


def test_preserve_overload_infeasible(tmp_path: Path) -> None:
    """When nothing fits: bounded infeasible, content/styles/count unchanged."""
    deck, ledger = _import_deck(["Перегруженный " * 400], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:1]
    tiny = {"content_box": {"x": 0.5, "y": 1.2, "w": 9.0, "h": 0.4},
            "unit": "in", "verified": True, "source": "test_template"}
    intent = _intent(slide.id, refs, candidate_id="over")
    tree = compile_composition(intent, slide, ledger, _rules(), tiny)
    res = resolve_layout(tree, slide, ledger, _rules(), tiny,
                         _bundle(), _measurer(), _options(tmp_path))
    assert res.status in ("infeasible", "needs_recompose",
                          "search_exhausted", "constraint_conflict")
    assert res.resolved_slide == {}
    assert any("needs_recompose" in u or "search_exhausted" in u
               or "constraint_conflict" in u or "CONSTRAINT_GEOMETRY_INFEASIBLE" in u
               for u in res.unresolved)
    # The refusal is bounded: the declared attempt budget is respected and the
    # solver's INFEASIBLE names the finite model it proved, never "too much
    # content" for the whole deck.
    assert len(res.attempts) <= 3
    assert any("CONSTRAINT_GEOMETRY_INFEASIBLE" in u for u in res.unresolved)


def test_global_attempt_budget(tmp_path: Path) -> None:
    """Per-call and cross-candidate budgets hold: a new candidate_id never
    resets the shared tracker."""
    deck, ledger = _import_deck(["Перегруженный " * 400], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:1]
    tiny = {"content_box": {"x": 0.5, "y": 1.2, "w": 9.0, "h": 0.4},
            "unit": "in", "verified": True, "source": "test_template"}
    tracker: dict = {"consumed": 0, "max_total": 2}
    for cid in ("budget-a", "budget-b", "budget-c"):
        intent = _intent(slide.id, refs, candidate_id=cid)
        tree = compile_composition(intent, slide, ledger, _rules(), tiny)
        res = resolve_layout(
            tree, slide, ledger, _rules(), tiny, _bundle(), _measurer(),
            _options(tmp_path, max_local_layout_attempts=5,
                     budget_tracker=tracker))
        for attempt in res.attempts:
            assert attempt.consumed_budget <= 5
    assert tracker["consumed"] <= 10
    intent = _intent(slide.id, refs, candidate_id="budget-d")
    tree = compile_composition(intent, slide, ledger, _rules(), tiny)
    res = resolve_layout(tree, slide, ledger, _rules(), tiny, _bundle(),
                         _measurer(),
                         _options(tmp_path, budget_tracker={"consumed": 7,
                                                            "max_total": 7}))
    assert res.status == "search_exhausted"
    assert any("BUDGET_EXHAUSTED" in u for u in res.unresolved)
    assert res.attempts == []


def test_resume_reuses_verified_results(tmp_path: Path) -> None:
    """Unchanged inputs hit the verified layout cache (no re-measurement);
    changed inputs recompute instead of reusing stale geometry."""
    deck, ledger = _import_deck(["Альфа", "Бета"], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:2]
    intent = _intent(slide.id, refs, candidate_id="resume")
    tree = compile_composition(intent, slide, ledger, _rules(), TEST_PROFILE)
    calls: list[str] = []

    def counting(request):
        calls.append(request.request_id)
        return _measurer() (request)

    first = resolve_layout(tree, slide, ledger, _rules(), TEST_PROFILE,
                           _bundle(), counting, _options(tmp_path))
    assert first.status == "resolved"
    n_calls = len(calls)
    assert n_calls > 0
    second = resolve_layout(tree, slide, ledger, _rules(), TEST_PROFILE,
                            _bundle(), counting, _options(tmp_path))
    assert len(calls) == n_calls
    assert (second.resolved_slide["geometry_emu"]
            == first.resolved_slide["geometry_emu"])
    moved = _intent(slide.id, refs, candidate_id="resume-moved",
                    zones=[{"zone_id": "z-half", "x_pct": 0.0, "y_pct": 0.0,
                            "w_pct": 50.0, "h_pct": 100.0, "role": "body",
                            "content_refs": list(refs)}])
    moved_tree = compile_composition(moved, slide, ledger, _rules(), TEST_PROFILE)
    third = resolve_layout(moved_tree, slide, ledger, _rules(), TEST_PROFILE,
                           _bundle(), counting, _options(tmp_path))
    assert len(calls) > n_calls
    assert third.status == "resolved"


def test_policy_calibration_invalidates_layout(tmp_path: Path) -> None:
    """The layout cache key covers calibration identity: geometry measured
    under another calibration (or changed spacing policy) is never reused."""
    import hashlib

    from slides_cli.gpn.layout import _layout_cache_key

    deck, ledger = _import_deck(["Альфа"], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:1]
    intent = _intent(slide.id, refs, candidate_id="cal")
    tree = compile_composition(intent, slide, ledger, _rules(), TEST_PROFILE)
    opts_a = _options(tmp_path, calibration_id="calibration-A")
    opts_b = _options(tmp_path, calibration_id="calibration-B")
    key_a = _layout_cache_key(tree, _rules(), TEST_PROFILE, _bundle(), opts_a)
    key_b = _layout_cache_key(tree, _rules(), TEST_PROFILE, _bundle(), opts_b)
    ha = hashlib.sha256(json.dumps(key_a, sort_keys=True, default=str).encode()).hexdigest()
    hb = hashlib.sha256(json.dumps(key_b, sort_keys=True, default=str).encode()).hexdigest()
    assert ha != hb


def test_generation_origin_preserved(tmp_path: Path) -> None:
    """The plan generation origin flows into the resolved slide (harness
    plans keep harness origin; nothing is silently re-attributed)."""
    deck, ledger = _import_deck(["Альфа"], tmp_path)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects if o.kind == "text"][:1]
    intent = _intent(slide.id, refs, candidate_id="origin")
    tree = compile_composition(intent, slide, ledger, _rules(), TEST_PROFILE)
    res = resolve_layout(tree, slide, ledger, _rules(), TEST_PROFILE,
                         _bundle(), _measurer(),
                         _options(tmp_path, generation_origin="harness_model"))
    assert res.status == "resolved"
    assert res.resolved_slide["generation_origin"] == "harness_model"


def _card_deck(tmp_path: Path, name: str = "card"):
    """One empty rounded-rect card + one textbox fully inside its bounds."""
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Emu

    from slides_cli.gpn.assets import AssetStore
    from slides_cli.gpn.importer import import_deck
    from slides_cli.gpn.provenance import build_ledger

    prs = fx.new_deck()
    slide = fx.blank_slide(prs)
    fx.add_textbox(slide, "Текст внутри карточки", left=2.5, top=2.5,
                   width=3.0, height=0.8, size=14)
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                           Emu(int(2 * 914400)), Emu(int(2 * 914400)),
                           Emu(int(4 * 914400)), Emu(int(2 * 914400)))
    src = tmp_path / f"{name}.pptx"
    prs.save(str(src))
    store = AssetStore(tmp_path / f"assets-{name}")
    deck, _ = import_deck(src, store)
    ledger = build_ledger(deck, decoration=[])
    return deck, ledger


def _is_empty_shape(o) -> bool:
    t = getattr(getattr(o, "payload", None), "text", None)
    paras = getattr(t, "paragraphs", None) or []
    return o.kind == "shape" and not any(
        "".join(getattr(r, "text", "") or ""
                for r in list(getattr(p, "runs", []) or []))
        for p in paras)


def test_overlap_relation_is_scoped(tmp_path: Path) -> None:
    """Container nesting is an exact-pair allowance; the same boxes
    without a declared pair refuse (no blanket allow-overlap)."""
    deck, ledger = _card_deck(tmp_path)
    slide = deck.slides[0]
    card = next(o.id for o in slide.objects if _is_empty_shape(o))
    text = next(o.id for o in slide.objects if o.kind == "text")
    zones = [
        {"zone_id": "z-card", "x_pct": 10.0, "y_pct": 20.0,
         "w_pct": 80.0, "h_pct": 40.0, "role": "body",
         "content_refs": [card]},
        {"zone_id": "z-text", "x_pct": 15.0, "y_pct": 25.0,
         "w_pct": 70.0, "h_pct": 30.0, "role": "body",
         "content_refs": [text]}]
    tiny = {"content_box": {"x": 0.5, "y": 1.2, "w": 9.0, "h": 4.3},
            "unit": "in", "verified": True, "source": "test_template"}
    allowed = _intent(slide.id, [card, text], candidate_id="nest-ok",
                      zones=zones, operators=["nest"],
                      allowed_overlaps=[[card, text]])
    tree = compile_composition(allowed, slide, ledger, _rules(), tiny)
    res = resolve_layout(tree, slide, ledger, _rules(), tiny,
                         _bundle(), _measurer(),
                         _options(tmp_path, layout_mode="recompose"))
    assert res.status == "resolved", res.unresolved

    denied = _intent(slide.id, [card, text], candidate_id="nest-no",
                     zones=zones, operators=["nest"])
    tree2 = compile_composition(denied, slide, ledger, _rules(), tiny)
    res2 = resolve_layout(tree2, slide, ledger, _rules(), tiny,
                          _bundle(), _measurer(),
                          _options(tmp_path, layout_mode="recompose"))
    # Without the pair the engine must move the text out (recorded), not
    # stack it inside silently.
    assert res2.status == "resolved", res2.unresolved
    moves = [m for a in res2.attempts for m in a.changes
             if isinstance(m, dict) and m.get("action") == "resolve_overlap_batch"]
    assert moves, "unpaired nesting must be repaired by an explicit move"

    # ...and when no room exists below, the same layout honestly refuses.
    flat = {"content_box": {"x": 0.5, "y": 1.2, "w": 9.0, "h": 0.4},
            "unit": "in", "verified": True, "source": "test_template"}
    tree3 = compile_composition(denied, slide, ledger, _rules(), flat)
    res3 = resolve_layout(tree3, slide, ledger, _rules(), flat,
                          _bundle(), _measurer(),
                          _options(tmp_path, layout_mode="recompose"))
    assert res3.status == "infeasible", res3.unresolved
    assert any("CONSTRAINT_GEOMETRY_INFEASIBLE" in u for u in res3.unresolved)
