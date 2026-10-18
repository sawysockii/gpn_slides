"""Real OOXML/native-render checks on explicit test-only fixtures.

These are engineering fixtures, never corporate design references.
"""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
import zipfile

import conftest as fx
import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

from slides_cli.gpn.assets import AssetStore
from slides_cli.gpn.constraint_layout import solve_measured_boxes
from slides_cli.gpn.importer import import_deck
from slides_cli.gpn.layout import _rich_text_paragraphs
from slides_cli.gpn.models import CompiledOntology, RoleStyle
from slides_cli.gpn.native import NativeDeckContext, emit_rich_text_block
from slides_cli.gpn.render import render_native_candidate

TEST_FACE = "Helvetica" if sys.platform == "darwin" else "DejaVu Sans"


def rules():
    return CompiledOntology(
        roles={
            "body": RoleStyle(
                role="body",
                allowed_faces=[TEST_FACE],
                default_size_pt=12,
                color_tokens=["text"],
            ),
            "caption": RoleStyle(
                role="caption",
                allowed_faces=[TEST_FACE],
                default_size_pt=10,
                color_tokens=["caption_color"],
            ),
        },
        colors={"text": "112233", "caption_color": "557799"},
    )


def test_rich_block_preserves_source_atom_ids_fields_breaks_links(tmp_path):
    source = fx.new_deck()
    slide = fx.blank_slide(source)
    fx.add_mixed_runs(slide)
    fx.add_fields_and_footnotes(slide)
    source_path = tmp_path / "test_only_source.pptx"
    source.save(source_path)
    store = AssetStore(tmp_path / "store")
    deck, ledger = import_deck(source_path, store)
    objects = {o.id: o for o in deck.slides[0].objects}
    refs = list(objects)
    paragraphs = _rich_text_paragraphs(objects, refs, {refs[-1]: "caption"})
    source_runs = {r.id for o in objects.values() for p in o.payload.paragraphs for r in p.runs}
    target_runs = {r["id"] for p in paragraphs for r in p["runs"]}
    assert source_runs <= target_runs
    target = fx.new_deck()
    out_slide = fx.blank_slide(target)
    context = NativeDeckContext(presentation=target)
    emit_rich_text_block(
        paragraphs,
        {"x": 12700 * 40, "y": 12700 * 50, "w": 12700 * 650, "h": 12700 * 400},
        out_slide,
        context,
        role="body",
        rules=rules(),
        node_id="merged",
    )
    output = tmp_path / "test_only_output.pptx"
    target.save(output)
    with zipfile.ZipFile(output) as archive:
        xml = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        fields = xml.findall(".//" + qn("a:fld"))
        assert {n.get("type") for n in fields} == {"slidenum", "datetime1"}
        assert len(xml.findall(".//" + qn("a:br"))) == 2
        assert xml.findall(".//" + qn("a:hlinkClick"))
        assert xml.findall(".//" + qn("a:rPr") + "[@baseline='30000']")
        assert xml.findall(".//" + qn("a:rPr") + "[@baseline='-25000']")
        assert {n.get("val") for n in xml.findall(".//" + qn("a:srgbClr"))} >= {"112233", "557799"}
        assert len(xml.findall(".//" + qn("p:sp"))) == 1
    reopened = Presentation(output)
    assert len(reopened.slides[0].shapes) == 1
    assert "05.10.2026" in reopened.slides[0].shapes[0].text
    actual, _ = import_deck(output, store)
    from slides_cli.gpn.export_pipeline import compare_export_preservation

    diff = compare_export_preservation(ledger, actual, context.output_map, {})
    assert diff.ok, diff.model_dump(mode="json")


def test_native_paragraph_spacing_has_real_units(tmp_path):
    deck = fx.new_deck()
    slide = fx.blank_slide(deck)
    emit_rich_text_block(
        [
            {
                "id": "p",
                "source_ref": "o",
                "space_before_pt": 4,
                "space_after_pt": 6,
                "line_spacing_mode": "points",
                "line_spacing_value": 16,
                "runs": [{"id": "r", "text": "Spacing"}],
            }
        ],
        {"x": 0, "y": 0, "w": 12700 * 300, "h": 12700 * 100},
        slide,
        NativeDeckContext(presentation=deck),
        role="body",
        rules=rules(),
    )
    p = slide.shapes[0].text_frame.paragraphs[0]
    assert p.space_before == Pt(4)
    assert p.space_after == Pt(6)
    assert p.line_spacing == Pt(16)


def test_fragmented_phrase_becomes_one_paragraph_with_exact_atom_ranges(tmp_path):
    source = fx.new_deck()
    slide = fx.blank_slide(source)
    for word in ("Одна", "смысловая", "фраза"):
        fx.add_textbox(slide, word)
    path = tmp_path / "fragmented.pptx"
    source.save(path)
    store = AssetStore(tmp_path / "store")
    deck, ledger = import_deck(path, store)
    objects = {o.id: o for o in deck.slides[0].objects}
    refs = list(objects)
    paragraphs = _rich_text_paragraphs(objects, refs, {}, [refs], " ")
    assert len(paragraphs) == 1
    assert "".join(r["text"] for r in paragraphs[0]["runs"]) == "Одна смысловая фраза"
    output = fx.new_deck()
    target = fx.blank_slide(output)
    context = NativeDeckContext(presentation=output)
    emit_rich_text_block(
        paragraphs,
        {"x": 0, "y": 0, "w": 12700 * 500, "h": 12700 * 40},
        target,
        context,
        role="body",
        rules=rules(),
        node_id="sentence",
    )
    final = tmp_path / "merged.pptx"
    output.save(final)
    assert len(Presentation(final).slides[0].shapes[0].text_frame.paragraphs) == 1
    actual, _ = import_deck(final, store)
    from slides_cli.gpn.export_pipeline import compare_export_preservation

    diff = compare_export_preservation(ledger, actual, context.output_map, {})
    assert diff.ok, diff.model_dump(mode="json")
    assert len(context.output_map) == len(ledger.atoms)
    # A missing middle word is detected against saved/reimported bytes.
    edited = Presentation(final)
    runs = edited.slides[0].shapes[0].text_frame.paragraphs[0].runs
    runs[2].text = ""
    edited.save(final)
    actual, _ = import_deck(final, store)
    assert not compare_export_preservation(ledger, actual, context.output_map, {}).ok


def test_resolved_scene_keeps_flow_and_roles_through_real_exporter(tmp_path):
    from types import SimpleNamespace

    from slides_cli.gpn.export_models import ResolvedSlide
    from slides_cli.gpn.export_pipeline import compare_export_preservation, emit_slide
    from slides_cli.gpn.layout import _build_resolved_slide
    from slides_cli.gpn.models import ComponentNode

    source = fx.new_deck()
    source_slide = fx.blank_slide(source)
    fx.add_textbox(source_slide, "01")
    fx.add_textbox(source_slide, "Подпись")
    path = tmp_path / "two_fragments.pptx"
    source.save(path)
    store = AssetStore(tmp_path / "store")
    deck, ledger = import_deck(path, store)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects]
    node = ComponentNode(
        id="rt-z",
        kind="rich_text",
        style_role="body",
        content_refs=refs,
        ref_roles={refs[0]: "body", refs[1]: "caption"},
        paragraph_groups=[refs],
        inline_separator=" ",
    )
    tree = SimpleNamespace(
        components={node.id: node}, slide_id=slide.id, candidate_id="c1", slide_background_token=""
    )
    resolved = ResolvedSlide.model_validate(
        _build_resolved_slide(
            tree,
            {},
            {node.id: {"x": 12700 * 40, "y": 12700 * 50, "w": 12700 * 500, "h": 12700 * 60}},
            False,
        )
    )
    target = fx.new_deck()
    target_slide = fx.blank_slide(target)
    context = NativeDeckContext(presentation=target)
    emit_slide(resolved, slide, context, target_slide, rules=rules())
    output = tmp_path / "actual_exporter.pptx"
    target.save(output)
    saved = Presentation(output)
    assert saved.slides[0].shapes[0].text == "01 Подпись"
    runs = saved.slides[0].shapes[0].text_frame.paragraphs[0].runs
    assert runs[0].font.size == Pt(12)
    assert runs[-1].font.size == Pt(10)
    assert str(runs[-1].font.color.rgb) == "557799"
    actual, _ = import_deck(output, store)
    diff = compare_export_preservation(ledger, actual, context.output_map, {})
    assert diff.ok, diff.model_dump(mode="json")


def test_real_layout_entry_accepts_strict_wire_projection_of_rich_spans(tmp_path):
    """Integration of source->composition->typed requests->real CP solver.

    Measurement is an explicit test-only protocol double here. Real native
    rendering is tested separately; this is not Chromium fit evidence.
    """
    from types import SimpleNamespace

    from slides_cli.gpn.composition import compile_composition
    from slides_cli.gpn.layout import resolve_layout
    from slides_cli.gpn.planner import LayoutIntent

    source = fx.new_deck()
    source_slide = fx.blank_slide(source)
    fx.add_textbox(source_slide, "01")
    fx.add_textbox(source_slide, "Подпись")
    path = tmp_path / "layout_projection.pptx"
    source.save(path)
    store = AssetStore(tmp_path / "store")
    deck, ledger = import_deck(path, store)
    slide = deck.slides[0]
    refs = [o.id for o in slide.objects]
    ontology = rules()
    ontology.canvas = slide.source_canvas
    intent = LayoutIntent.model_validate(
        {
            "slide_id": slide.id,
            "candidate_id": "c1",
            "content_refs": refs,
            "reading_order": refs,
            "dominant_ref": refs[0],
            "zones": [
                {
                    "zone_id": "z",
                    "x_pct": 0,
                    "y_pct": 0,
                    "w_pct": 100,
                    "h_pct": 100,
                    "role": "body",
                    "content_refs": refs,
                    "ref_roles": {refs[-1]: "caption"},
                    "paragraph_groups": [refs],
                    "inline_separator": " ",
                }
            ],
        }
    )
    profile = {
        "content_box": {"x": 0.5, "y": 0.7, "w": 9, "h": 4.5},
        "unit": "in",
        "verified": True,
        "source": "test_only",
    }
    tree = compile_composition(intent, slide, ledger, ontology, profile)
    bundle = SimpleNamespace(
        bindings=[
            SimpleNamespace(
                role_or_face=TEST_FACE,
                browser_alias="MEASURE_FACE_012345abcdef",
                path="/test-only/protocol-double",
                sha256="0" * 64,
                face_index=None,
            )
        ],
        missing_faces=[],
        bundle_hash="test-only",
    )
    requests = []

    def test_only_measurer(request):
        requests.append(request)
        assert request.profile_kind == "test_only"
        return {
            "items": [
                {
                    "component_id": comp.id,
                    "required_outer_height_pt": 24,
                    "occupied_height_pt": 16,
                    "intrinsic_min_width_pt": 30,
                    "intrinsic_preferred_width_pt": 100,
                    "overflow_x": False,
                    "overflow_y": False,
                    "confidence": "test_only",
                }
                for comp in request.components
            ]
        }

    counters = {}
    result = resolve_layout(
        tree,
        slide,
        ledger,
        ontology,
        profile,
        bundle,
        test_only_measurer,
        {
            "profile_kind": "test_only",
            "use_cache": False,
            "layout_mode": "recompose",
            "runtime_counters": counters,
        },
    )
    assert result.status == "resolved", result.unresolved
    assert requests
    assert counters["solver_calls"] >= 1
    paragraphs = requests[0].components[0].paragraphs
    assert len(paragraphs) == 1
    assert paragraphs[0].runs[-1].style.size_pt == 10


def test_joint_solver_native_saved_boxes_and_edit_roundtrip(tmp_path):
    hints = {f"b{i}": {"x": 40, "y": 50, "w": 180, "h": 40} for i in range(9)}
    result = solve_measured_boxes(hints, {"x": 40, "y": 50, "w": 600, "h": 300}, seconds=2)
    assert result.status in {"OPTIMAL", "FEASIBLE"}
    deck = fx.new_deck()
    slide = fx.blank_slide(deck)
    context = NativeDeckContext(presentation=deck)
    for key, box in result.boxes.items():
        emit_rich_text_block(
            [
                {
                    "id": key,
                    "source_ref": key,
                    "runs": [{"id": key + ":r", "text": f"Проверка {key}"}],
                }
            ],
            {k: round(v * 12700) for k, v in box.items()},
            slide,
            context,
            role="body",
            rules=rules(),
        )
    path = tmp_path / "test_only_nine_blocks.pptx"
    deck.save(path)
    reopened = Presentation(path)
    actual = [
        {"x": s.left / 12700, "y": s.top / 12700, "w": s.width / 12700, "h": s.height / 12700}
        for s in reopened.slides[0].shapes
    ]
    assert actual == list(result.boxes.values())
    reopened.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = "Изменено"
    edited = tmp_path / "test_only_edited.pptx"
    reopened.save(edited)
    assert Presentation(edited).slides[0].shapes[0].text == "Изменено"


@pytest.mark.skipif(
    shutil.which("soffice") is None or shutil.which("pdftoppm") is None,
    reason="real native renderer unavailable",
)
def test_actual_native_render_includes_hidden_slide(tmp_path):
    deck = fx.new_deck()
    for index in range(2):
        slide = fx.blank_slide(deck)
        if index:
            slide._element.set("show", "0")
        emit_rich_text_block(
            [
                {
                    "id": f"p{index}",
                    "source_ref": f"o{index}",
                    "runs": [{"id": f"r{index}", "text": f"Native probe {index + 1}"}],
                }
            ],
            {"x": 12700 * 40, "y": 12700 * 50, "w": 12700 * 400, "h": 12700 * 60},
            slide,
            NativeDeckContext(presentation=deck),
            role="body",
            rules=rules(),
        )
    path = tmp_path / "test_only_hidden.pptx"
    deck.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    probe = render_native_candidate(
        path,
        {
            "output_dir": str(tmp_path / "render"),
            "expected_faces": [TEST_FACE],
            "expected_page_count": 2,
            "dpi": 72,
        },
        None,
    )
    assert probe.rendered, probe.issues
    assert probe.page_count == 2, probe.issues
    assert probe.result == "pass", probe.issues
    assert len(probe.page_images) == 2
    assert before == hashlib.sha256(path.read_bytes()).hexdigest()
    assert Presentation(path).slides[1]._element.get("show") == "0"
    (tmp_path / "native_probe.json").write_text(json.dumps(probe.model_dump(mode="json"), indent=2))
